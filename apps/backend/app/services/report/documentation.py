"""Org documentation workbook (roadmap item 11, part a): the org's
configuration as an Excel deliverable, from the last sync, with who can
reach each thing documented. Config Workbook parity rows #1, #2, #3 (in
part), #5, #10, #12 and #17 (docs/config-workbook-parity.md); validation
rules, tabs, Visualforce and the other new reads follow in part b.

Sheets: About, Objects, Fields, Picklist Values, Object Access, Field
Access, Assignments, Profile Users, Record Types, Layouts, Unused Layouts,
Record Pages, Flows, Apex Classes; part b adds the rest of SHEETS. Part c:
Fields on Layouts (#4), Flow Elements and Flow Fields (#9), Approval Steps
(#18), External Data Sources and record types through permission sets (#11).
Part c2: Custom Metadata Types and, for the types an admin turned on,
Custom Metadata Records (#7).

Configuration sheets work at every privacy level. User counts and the
per-user sheets (Assignments, Profile Users) need user records, so at
Metadata only they're left out and the About sheet says so. Names come
from the database, which privacy mode has already masked; no emails.
"""
from __future__ import annotations

import io
import re
from collections import defaultdict
from datetime import datetime, timezone
from typing import Any, Dict, List, Optional, Set, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.domain.models import (
    FieldPermissionSnapshot,
    ObjectPermissionSnapshot,
    Organization,
    PermissionSetAssignmentSnapshot,
    PermissionSetGroupComponentSnapshot,
    PermissionSetGroupSnapshot,
    PermissionSetSnapshot,
    ProfileSnapshot,
    UserSnapshot,
)
from app.services import privacy_mode
from app.services import custom_metadata as cmdt
from app.services import metadata_detail as md_detail
from app.services import schema_metadata as sm
from app.services import ui_access as ua
from app.services.org_docs import latest_org_docs
from app.services.effective_access import EffectiveAccessService, _Component, _evaluate, _Source
from app.services.report.workbook import Column, _safe, _table

OBJECT_FLAGS6 = ("read", "create", "edit", "delete", "view_all", "modify_all")
OBJECT_LETTERS = {"read": "R", "create": "C", "edit": "E", "delete": "D", "view_all": "VA", "modify_all": "MA"}
NOT_AT_METADATA_ONLY = "Not available at this org's privacy level (metadata only)."
NOT_SYNCED_SCHEMA = "Field metadata isn't synced yet: sync the org to read field types, formulas and relationships."
NOT_SYNCED_UI = "Not read yet: sync the org to read page layouts, record pages, Apex and flows."
SHEETS = ("About", "Objects", "Fields", "Picklist Values", "Picklists by Record Type", "Validation Rules",
          "Object Access", "Field Access", "Tab Visibility", "Assignments", "Profile Users", "Login Restrictions",
          "Record Types", "Layouts", "Fields on Layouts", "Unused Layouts", "Record Pages", "Buttons and Links",
          "List Views", "Field Sets", "Apps", "Flows", "Flow Elements", "Flow Fields", "Approval Processes",
          "Approval Steps", "Apex Classes", "Apex Triggers", "Visualforce", "Lightning Components", "Static Resources",
          "Custom Permissions", "External Data Sources", "Custom Metadata Types", "Custom Metadata Records")
NOT_SYNCED_DOCS = "Not read yet: sync the org to read this part of its configuration."
DAY_ABBR = {"monday": "Mon", "tuesday": "Tue", "wednesday": "Wed", "thursday": "Thu", "friday": "Fri",
            "saturday": "Sat", "sunday": "Sun"}


def _when(iso: Optional[str]) -> str:
    if not iso:
        return "Not synced yet"
    try:
        return datetime.fromisoformat(iso.replace("Z", "+00:00")).astimezone(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    except ValueError:
        return iso


def _yes(v: Optional[bool]) -> str:
    return "" if v is None else ("Yes" if v else "No")


class _Org:
    """Everything the sheets read, loaded once."""

    def __init__(self) -> None:
        self.profiles: Dict[str, str] = {}
        self.ps_rows: Dict[str, Any] = {}
        self.groups: Dict[str, Any] = {}
        self.components: List[Any] = []
        self.components_by_group: Dict[str, List[str]] = defaultdict(list)
        self.assignments: List[Any] = []
        self.users: Dict[str, Any] = {}
        self.object_rows: List[Any] = []
        self.field_rows: List[Any] = []
        self.schema: Optional[Dict[str, Any]] = None
        self.ui: Optional[Dict[str, Any]] = None
        self.docs: Optional[Dict[str, Any]] = None
        self.access: Dict[str, Dict[str, Dict[str, Dict[str, bool]]]] = {}
        self.users_allowed = False

    @property
    def group_names(self) -> Dict[str, str]:
        return {gid: g.master_label or g.developer_name or gid for gid, g in self.groups.items()}

    def columns(self) -> List[Tuple[str, str, _Source]]:
        """The access matrices' columns: (kind, heading, source) for each
        profile, permission set and permission set group, by name."""
        out = []
        for ps in self.ps_rows.values():
            if ps.is_owned_by_profile and ps.profile_id:
                name = self.profiles.get(ps.profile_id, ps.profile_id)
                out.append(("profile", f"{name} (Profile)", _Source("profile", name, [_Component(ps.salesforce_id, name)])))
            elif (ps.ps_type or "Regular") not in ("Group", "Muting") and not ps.is_muting:
                label = ps.label or ps.name
                out.append(("permission_set", label, _Source("permission_set", label, [_Component(ps.salesforce_id, label)])))
        for gid, name in self.group_names.items():
            comps = []
            for cid in self.components_by_group.get(gid, []):
                c = self.ps_rows.get(cid)
                comps.append(_Component(cid, (c.label or c.name) if c else cid, bool(c and (c.is_muting or c.ps_type == "Muting"))))
            out.append(("permission_set_group", f"{name} (Group)", _Source("permission_set_group", name, comps, group_id=gid)))
        order = {"profile": 0, "permission_set": 1, "permission_set_group": 2}
        return sorted(out, key=lambda c: (order[c[0]], c[1].lower()))


async def _load(db: AsyncSession, org: Organization) -> _Org:
    from app.api.routes.ui_access import latest_ui_access

    o = _Org()
    oid = org.id
    o.users_allowed = privacy_mode.allows(org, privacy_mode.USER_RECORDS)
    o.profiles = {p.salesforce_id: p.name for p in (await db.execute(
        select(ProfileSnapshot).where(ProfileSnapshot.organization_id == oid))).scalars().all()}
    o.ps_rows = {p.salesforce_id: p for p in (await db.execute(
        select(PermissionSetSnapshot).where(PermissionSetSnapshot.organization_id == oid))).scalars().all()}
    o.groups = {g.salesforce_id: g for g in (await db.execute(
        select(PermissionSetGroupSnapshot).where(PermissionSetGroupSnapshot.organization_id == oid))).scalars().all()}
    o.components = list((await db.execute(select(PermissionSetGroupComponentSnapshot).where(
        PermissionSetGroupComponentSnapshot.organization_id == oid))).scalars().all())
    for c in o.components:
        o.components_by_group[c.permission_set_group_id].append(c.permission_set_id)
    o.object_rows = list((await db.execute(select(ObjectPermissionSnapshot).where(
        ObjectPermissionSnapshot.organization_id == oid))).scalars().all())
    o.field_rows = list((await db.execute(select(FieldPermissionSnapshot).where(
        FieldPermissionSnapshot.organization_id == oid))).scalars().all())
    o.schema = await sm.latest_schema(db, oid)
    o.ui = await latest_ui_access(db, oid)
    o.docs = await latest_org_docs(db, oid)
    if o.users_allowed:
        o.assignments = list((await db.execute(select(PermissionSetAssignmentSnapshot).where(
            PermissionSetAssignmentSnapshot.organization_id == oid))).scalars().all())
        o.users = {u.salesforce_id: u for u in (await db.execute(select(UserSnapshot).where(
            UserSnapshot.organization_id == oid, UserSnapshot.is_active.is_(True)))).scalars().all()}
        o.access = await EffectiveAccessService(db).org_access(oid)
    return o


def _object_perms6(o: _Org) -> Dict[str, Dict[str, Dict[str, bool]]]:
    out: Dict[str, Dict[str, Dict[str, bool]]] = defaultdict(dict)
    for p in o.object_rows:
        out[p.parent_id][p.sobject_type] = {
            "read": bool(p.permissions_read), "create": bool(p.permissions_create), "edit": bool(p.permissions_edit),
            "delete": bool(p.permissions_delete), "view_all": bool(p.permissions_view_all_records),
            "modify_all": bool(p.permissions_modify_all_records)}
    return out


def _field_perms(o: _Org) -> Dict[str, Dict[str, Dict[str, bool]]]:
    out: Dict[str, Dict[str, Dict[str, bool]]] = defaultdict(dict)
    for p in o.field_rows:
        cur = out[p.parent_id].get(p.field, {"read": False, "edit": False})
        out[p.parent_id][p.field] = {"read": cur["read"] or bool(p.permissions_read),
                                     "edit": cur["edit"] or bool(p.permissions_edit)}
    return out


def _user_counts(o: _Org, kind: str) -> Dict[str, Dict[str, int]]:
    """key (object or Object.Field) -> flag -> active users holding it."""
    out: Dict[str, Dict[str, int]] = defaultdict(lambda: defaultdict(int))
    for acc in o.access.values():
        for key, flags in acc[kind].items():
            for f, on in flags.items():
                if on:
                    out[key][f] += 1
    return out


def _described_objects(o: _Org) -> Dict[str, Dict[str, Any]]:
    return (o.schema or {}).get("objects") or {}


def _all_objects(o: _Org) -> List[str]:
    return sorted({p.sobject_type for p in o.object_rows} | set(_described_objects(o)))


# --------------------------------------------------------------------------- sheets

def _about(ws: Worksheet, org: Organization, o: _Org, now: datetime) -> None:
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 100
    level = {privacy_mode.FULL: "Full", privacy_mode.MASKED: "Masked users",
             privacy_mode.METADATA_ONLY: "Metadata only"}.get(privacy_mode.mode_of(org), privacy_mode.mode_of(org))
    rows: List[Tuple[str, Any]] = []
    if org.is_demo:
        rows.append(("Demo data", "This org's data is synthetic."))
    rows += [
        ("Organization", org.name),
        ("Generated", now.strftime("%Y-%m-%d %H:%M UTC")),
        ("Privacy level", level or ""),
        ("Field metadata read", _when((o.schema or {}).get("collected_at"))),
        ("Layouts, pages, Apex, flows read", _when((o.ui or {}).get("collected_at"))),
        ("Who can reach it", "Counts are active users, by the same rules as Newton's user pages: profile, permission "
                             "sets, and permission set groups with muting applied." if o.users_allowed else NOT_AT_METADATA_ONLY
         + " User counts and the Assignments and Profile Users sheets are left out."),
        ("Access letters", "R Read, C Create, E Edit, D Delete, VA View All, MA Modify All (Object Access); "
                           "Read or Read/Edit (Field Access)."),
        ("Field access", "Fields whose field-level security can't be set (record id, required fields, system "
                         "fields) follow the object's access; the Fields sheet says which."),
        ("Tab visibility", "Default On or Default Off per profile and permission set; a blank cell means no "
                           "tab setting was returned for it."),
        ("Documentation read", _when((o.docs or {}).get("collected_at"))),
        ("Fields on layouts", "Who gets a layout comes from the profile layout assignments; a field on it shows only "
                              "for those who can read the field, which the last column counts."),
        ("Custom metadata", "Records are read only for the types an admin turned on for this org (Privacy & Data); "
                            "the Custom Metadata Types sheet says which."),
        ("Not included yet", "Approval process entry criteria and approvers (they need the Metadata API, which isn't "
                             "confirmed to work with Newton's permissions)."),
    ]
    for note in ((o.schema or {}).get("notes") or []) + ((o.ui or {}).get("notes") or []) + ((o.docs or {}).get("notes") or []):
        rows.append(("Note", note))
    for i, (k, v) in enumerate(rows, start=1):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        cell = ws.cell(row=i, column=2, value=_safe(v))
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def _objects(ws: Worksheet, o: _Org) -> None:
    described = _described_objects(o)
    counts = _user_counts(o, "objects") if o.users_allowed else {}
    perms = _object_perms6(o)
    grants: Dict[str, Dict[str, Set[str]]] = defaultdict(lambda: defaultdict(set))
    for kind, heading, src in o.columns():
        if kind == "permission_set_group":
            continue  # a group's sets are counted as sets
        for obj, flags in perms.get(src.components[0].ps_id, {}).items():
            for f in ("read", "view_all", "modify_all"):
                if flags[f]:
                    grants[obj][f].add(heading)
    cols = [Column("Object", 28), Column("Label", 24), Column("Custom", 9), Column("Fields", 9),
            Column("Master-detail parent", 24), Column("Profiles and sets with Read", 14),
            Column("View All granted by", 36, wrap=True), Column("Modify All granted by", 36, wrap=True)]
    if o.users_allowed:
        cols += [Column(f"Users with {f.title()}", 12) for f in ("read", "create", "edit", "delete")]
    rows = []
    for name in _all_objects(o):
        d = described.get(name) or {}
        md = sorted({ref for f in (d.get("fields") or {}).values() if f.get("relationship") == "master_detail"
                     for ref in f.get("referenceTo") or []})
        row = [name, d.get("label") or "", _yes(name.endswith("__c")), len(d.get("fields") or {}) if d else None,
               ", ".join(md), len(grants[name]["read"]), ", ".join(sorted(grants[name]["view_all"])),
               ", ".join(sorted(grants[name]["modify_all"]))]
        if o.users_allowed:
            row += [counts.get(name, {}).get(f, 0) for f in ("read", "create", "edit", "delete")]
        rows.append(row)
    _table(ws, cols, rows, "No objects synced.")


def _fields(ws: Worksheet, o: _Org) -> None:
    described = _described_objects(o)
    if not described:
        _table(ws, [Column("Object", 28)], [], NOT_SYNCED_SCHEMA)
        return
    fcounts = _user_counts(o, "fields") if o.users_allowed else {}
    ocounts = _user_counts(o, "objects") if o.users_allowed else {}
    cols = [Column("Object", 26), Column("Field", 28), Column("Label", 26), Column("Type", 26), Column("Length", 9),
            Column("Required", 10), Column("Formula", 50, wrap=True), Column("References", 22),
            Column("Encryption", 12), Column("External ID", 11), Column("Unique", 9), Column("Help text", 30, wrap=True),
            Column("Access follows", 22), Column("Source", 16)]
    if o.users_allowed:
        cols += [Column("Users with Read", 12), Column("Users with Edit", 12)]
    rows = []
    for obj in sorted(described):
        for name, f in described[obj]["fields"].items():
            hidden = f.get("source") == "field_definition"
            follows_object = f.get("permissionable") is False
            row = [obj, name, f.get("label") or "", sm.type_label(f), f.get("length"),
                   "Not known" if hidden else _yes(f.get("required")),
                   f.get("calculatedFormula") or ("Formula; text not known" if f.get("formula") else ""),
                   ", ".join(f.get("referenceTo") or []),
                   "Not known" if hidden else {"shield": "Shield", "classic": "Classic"}.get(f.get("encryption") or "", "None"),
                   _yes(f.get("externalId")) if not hidden else "", _yes(f.get("unique")) if not hidden else "",
                   f.get("inlineHelpText") or "",
                   "Object access" if follows_object else "Field-level security",
                   "FieldDefinition (hidden from the integration user)" if hidden else "Describe"]
            if o.users_allowed:
                if follows_object:
                    oc = ocounts.get(obj, {})
                    row += [oc.get("read", 0), oc.get("edit", 0)]
                else:
                    fc = fcounts.get(f"{obj}.{name}", {})
                    row += [fc.get("read", 0), fc.get("edit", 0)]
            rows.append(row)
    _table(ws, cols, rows, "No fields described.")


def _picklists(ws: Worksheet, o: _Org) -> None:
    described = _described_objects(o)
    cols = [Column("Object", 26), Column("Field", 28), Column("Value", 30), Column("Label", 30),
            Column("Active", 9), Column("Default", 9)]
    rows = [[obj, name, v.get("value"), v.get("label"), _yes(v.get("active", True)), _yes(bool(v.get("default")))]
            for obj in sorted(described) for name, f in described[obj]["fields"].items()
            for v in f.get("picklistValues") or []]
    _table(ws, cols, rows, NOT_SYNCED_SCHEMA if not described else
           "No picklist values in the synced describes (hidden fields have none).")


def _matrix(ws: Worksheet, o: _Org, kind: str) -> None:
    """Rows: objects (or fields); columns: every profile, permission set and group."""
    columns = o.columns()
    flags: Tuple[str, ...]
    if kind == "objects":
        perms, flags = _object_perms6(o), OBJECT_FLAGS6
        keys = _all_objects(o)
    else:
        perms, flags = _field_perms(o), ("read", "edit")
        keys = sorted({p.field for p in o.field_rows})
    evaluated = [_evaluate(src, perms, flags).granted for _, _, src in columns]
    head = [Column("Object" if kind == "objects" else "Field", 34)] + [Column(h, 14) for _, h, _ in columns]

    def cell(g: Dict[str, Dict[str, bool]], key: str) -> str:
        p = g.get(key)
        if not p:
            return ""
        if kind == "objects":
            return " ".join(OBJECT_LETTERS[f] for f in flags if p.get(f))
        return "Read/Edit" if p.get("edit") else ("Read" if p.get("read") else "")

    rows = [[key] + [cell(g, key) for g in evaluated] for key in keys]
    _table(ws, head, rows, "No permissions synced.")
    for i in range(2, len(columns) + 2):
        ws.cell(row=1, column=i).alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"


def _assignments(ws: Worksheet, o: _Org) -> None:
    cols = [Column("Permission set or group", 34), Column("Type", 22), Column("User", 26), Column("Username", 34)]
    if not o.users_allowed:
        _table(ws, cols, [], NOT_AT_METADATA_ONLY)
        return
    rows = []
    for a in o.assignments:
        u = o.users.get(a.assignee_id)
        ps = o.ps_rows.get(a.permission_set_id)
        if u is None or (ps is not None and ps.is_owned_by_profile):
            continue
        if a.permission_set_id in o.groups:
            name, kind = o.group_names[a.permission_set_id], "Permission set group"
        elif ps is not None and (ps.ps_type or "") == "Group":
            gid = (a.raw_data or {}).get("PermissionSetGroupId") or (ps.raw_data or {}).get("PermissionSetGroupId")
            name, kind = o.group_names.get(gid or "", ps.label or ps.name), "Permission set group"
        else:
            name, kind = (ps.label or ps.name) if ps else a.permission_set_id, "Permission set"
        rows.append([name, kind, u.name, u.username])
    rows.sort(key=lambda r: (r[0].lower(), r[2].lower()))
    _table(ws, cols, rows, "No permission set assignments synced.")


def _profile_users(ws: Worksheet, o: _Org) -> None:
    cols = [Column("Profile", 28), Column("User", 26), Column("Username", 34), Column("Permission sets and groups", 60, wrap=True)]
    if not o.users_allowed:
        _table(ws, cols, [], NOT_AT_METADATA_ONLY)
        return
    held: Dict[str, List[str]] = defaultdict(list)
    for a in o.assignments:
        ps = o.ps_rows.get(a.permission_set_id)
        if ps is not None and ps.is_owned_by_profile:
            continue
        name = o.group_names.get(a.permission_set_id) or ((ps.label or ps.name) if ps else a.permission_set_id)
        if name not in held[a.assignee_id]:
            held[a.assignee_id].append(name)
    rows = sorted(([o.profiles.get(u.profile_id, u.profile_id or ""), u.name, u.username, ", ".join(sorted(held[uid]))]
                   for uid, u in o.users.items()), key=lambda r: (r[0].lower(), r[1].lower()))
    _table(ws, cols, rows, "No active users synced.")


def _record_types(ws: Worksheet, o: _Org) -> None:
    cols = [Column("Object", 24), Column("Record type", 26), Column("API name", 26), Column("Active", 9),
            Column("Available to profiles", 50, wrap=True), Column("Default for profiles", 40, wrap=True),
            Column("Available through permission sets", 44, wrap=True)]
    if not o.ui:
        _table(ws, cols, [], NOT_SYNCED_UI)
        return
    vis = o.ui.get("record_type_visibility")
    rows = []
    for rt in sorted(o.ui.get("record_types", []), key=lambda r: (r["object"], r["name"])):
        key = f"{rt['object']}.{rt['developer_name']}"
        if vis is None:
            avail = default = "Not read"
        else:
            avail = ", ".join(sorted(o.profiles.get(pid, pid) for pid, objs in vis.items()
                                     if key in (objs.get(rt["object"]) or {}).get("available", [])))
            default = ", ".join(sorted(o.profiles.get(pid, pid) for pid, objs in vis.items()
                                       if (objs.get(rt["object"]) or {}).get("default") == key))
        set_rts = o.ui.get("set_record_types")
        through = "Not read" if set_rts is None else ", ".join(sorted(
            _set_label(o, ps_id) for ps_id, objs in set_rts.items() if key in (objs.get(rt["object"]) or [])))
        rows.append([rt["object"], rt["name"], rt["developer_name"], _yes(rt.get("active", True)), avail, default, through])
    _table(ws, cols, rows, "No record types.")


def _set_label(o: _Org, ps_id: str) -> str:
    """A permission set's label, with the groups it sits in."""
    ps = o.ps_rows.get(ps_id)
    name = (ps.label or ps.name) if ps else ps_id
    groups = sorted(o.group_names.get(gid, gid) for gid, comps in o.components_by_group.items() if ps_id in comps)
    return name + (f" (in {', '.join(groups)})" if groups else "")


def _active_by_profile(o: _Org) -> Dict[str, int]:
    out: Dict[str, int] = defaultdict(int)
    for u in o.users.values():
        out[u.profile_id or ""] += 1
    return out


def _layouts(ws: Worksheet, o: _Org) -> None:
    cols = [Column("Object", 24), Column("Layout", 36), Column("Profile", 28), Column("Record type", 26)]
    if o.users_allowed:
        cols.append(Column("Active users with the profile", 14))
    if not o.ui:
        _table(ws, cols, [], NOT_SYNCED_UI)
        return
    layouts = {lay["id"]: lay for lay in o.ui.get("layouts", [])}
    rts = {r["id"]: r for r in o.ui.get("record_types", [])}
    per_profile = _active_by_profile(o)
    rows = []
    for pl in o.ui.get("profile_layouts", []):
        rt = rts.get(pl.get("record_type_id") or "")
        row = [pl["object"], (layouts.get(pl["layout_id"]) or {}).get("name", pl["layout_id"]),
               o.profiles.get(pl["profile_id"], pl["profile_id"]), rt["name"] if rt else "Master (no record type)"]
        if o.users_allowed:
            row.append(per_profile.get(pl["profile_id"], 0))
        rows.append(row)
    rows.sort(key=lambda r: (r[0], r[1].lower(), r[2].lower(), r[3]))
    _table(ws, cols, rows, "No layout assignments.")


def _layout_fields(ws: Worksheet, o: _Org) -> None:
    """The field x layout matrix (parity #4): every field on every layout, the
    profiles that get the layout, and how many of their active users can read
    the field (a field on a layout shows only to those who can read it)."""
    cols = [Column("Object", 22), Column("Field", 28), Column("Label", 24), Column("Layout", 34), Column("Section", 22),
            Column("Behavior", 11), Column("Layout assigned to profiles", 44, wrap=True)]
    if o.users_allowed:
        cols += [Column("Active users who get the layout", 14), Column("Of them, can read the field", 14)]
    if not o.ui:
        _table(ws, cols, [], NOT_SYNCED_UI)
        return
    placed = o.ui.get("layout_fields")
    if placed is None:
        _table(ws, cols, [], "The fields on each layout weren't read: sync the org again to read them.")
        return
    described = _described_objects(o)
    profiles_of: Dict[str, Set[str]] = defaultdict(set)
    for pl in o.ui.get("profile_layouts", []):
        profiles_of[pl["layout_id"]].add(pl["profile_id"])
    users_by_profile: Dict[str, List[str]] = defaultdict(list)
    for uid, u in o.users.items():
        users_by_profile[u.profile_id or ""].append(uid)
    rows = []
    for lay in sorted(o.ui.get("layouts", []), key=lambda x: ((x.get("object") or ""), x["name"].lower())):
        obj = lay.get("object") or ""
        profs = profiles_of.get(lay["id"], set())
        users = [uid for pid in profs for uid in users_by_profile.get(pid, [])]
        for item in placed.get(lay["id"], []):
            meta = ((described.get(obj) or {}).get("fields") or {}).get(item["field"]) or {}
            row = [obj, item["field"], meta.get("label") or "", lay["name"], item.get("section") or "",
                   item.get("behavior") or "", ", ".join(sorted(o.profiles.get(p, p) for p in profs)) or "Not assigned"]
            if o.users_allowed:
                follows_object = meta.get("permissionable") is False
                key, kind = (obj, "objects") if follows_object else (f"{obj}.{item['field']}", "fields")
                row += [len(users), sum(1 for uid in users if ((o.access.get(uid) or {}).get(kind, {}).get(key) or {}).get("read"))]
            rows.append(row)
    _table(ws, cols, rows, "No fields read from any layout.")


def _unused_layouts(ws: Worksheet, o: _Org) -> None:
    cols = [Column("Object", 24), Column("Layout", 40), Column("Why it's listed", 70, wrap=True)]
    if not o.ui:
        _table(ws, cols, [], NOT_SYNCED_UI)
        return
    used = {pl["layout_id"] for pl in o.ui.get("profile_layouts", [])}
    rows = [[lay["object"], lay["name"], "No profile or record type is assigned this layout in the synced layout "
             "assignments (ProfileLayout)."]
            for lay in sorted(o.ui.get("layouts", []), key=lambda x: (x["object"], x["name"].lower())) if lay["id"] not in used]
    _table(ws, cols, rows, "Every layout is assigned to at least one profile.")


def _record_pages(ws: Worksheet, o: _Org) -> None:
    cols = [Column("Page", 34), Column("API name", 30), Column("Object", 22), Column("Assigned as", 70, wrap=True)]
    if not o.ui:
        _table(ws, cols, [], NOT_SYNCED_UI)
        return
    uses: Dict[str, List[str]] = defaultdict(list)
    for obj, by_ff in (o.ui.get("org_default_pages") or {}).items():
        for ff, page in by_ff.items():
            uses[page].append(f"Org default ({'Desktop' if ff == 'Large' else 'Phone'})")
    for app in o.ui.get("apps", []):
        for ov in app.get("overrides", []):
            what = [f"App: {app['label']}"]
            if ov.get("record_type"):
                what.append(f"record type {ov['record_type'].split('.', 1)[-1]}")
            if ov.get("profile"):
                what.append(f"profile {ov['profile']}")
            uses[ov["page"]].append(", ".join(what) + f" ({'Desktop' if ov.get('form_factor') == 'Large' else 'Phone'})")
    rows = [[p.get("label") or p["developer_name"], p["developer_name"], p.get("object") or "",
             "; ".join(uses.get(p["developer_name"], [])) or "Not assigned"]
            for p in sorted(o.ui.get("flexipages", []), key=lambda p: ((p.get("object") or ""), (p.get("label") or "").lower()))]
    _table(ws, cols, rows, "No Lightning record pages.")


def _grant_names(o: _Org, kind: str) -> Dict[str, List[str]]:
    """Entity id -> the profiles, sets (and the groups they sit in) that grant it."""
    in_groups: Dict[str, List[str]] = defaultdict(list)
    for gid, comps in o.components_by_group.items():
        for cid in comps:
            in_groups[cid].append(o.group_names.get(gid, gid))
    out: Dict[str, List[str]] = defaultdict(list)
    for parent, kinds in ((o.ui or {}).get("setup_entity_access") or {}).items():
        ps = o.ps_rows.get(parent)
        if ps is not None and ps.is_owned_by_profile:
            name = f"{o.profiles.get(ps.profile_id, ps.profile_id)} (Profile)"
        else:
            name = (ps.label or ps.name) if ps else parent
            if in_groups.get(parent):
                name += f" (in {', '.join(sorted(in_groups[parent]))})"
        for eid in (kinds or {}).get(kind, []):
            out[eid].append(name)
    return out


GRANT_KINDS = ("ApexClass", "ApexPage", "CustomPermission", "TabSet", "ExternalDataSource")


def _per_user_runs(o: _Org) -> Tuple[Dict[str, int], Dict[str, Dict[str, int]]]:
    """Flow id -> active users who can run it; for each SetupEntityAccess kind
    (Apex class, Visualforce page, custom permission, app), entity id -> active
    users granted it, through their profile, sets and groups."""
    flows: Dict[str, int] = defaultdict(int)
    granted: Dict[str, Dict[str, int]] = {k: defaultdict(int) for k in GRANT_KINDS}
    if not (o.ui and o.users_allowed):
        return flows, granted
    by_user: Dict[str, List[Any]] = defaultdict(list)
    for a in o.assignments:
        by_user[a.assignee_id].append(a)
    for uid, u in o.users.items():
        sources = ua.user_sources(u, o.profiles, o.ps_rows, by_user.get(uid, []), o.components, o.group_names)
        for f in ua.resolve_flows(o.ui, sources, uid)["flows"]:
            if f["can_run"] is True:
                flows[f["id"]] += 1
        for kind in GRANT_KINDS:
            for eid in ua._granted(o.ui, sources, kind):
                granted[kind][eid] += 1
    return flows, granted


def _flows(ws: Worksheet, o: _Org, runs: Dict[str, int]) -> None:
    cols = [Column("Flow", 34), Column("API name", 32), Column("Type", 20), Column("Trigger", 18), Column("Active", 9),
            Column("Requires explicit access", 14), Column("Explicit access granted by", 50, wrap=True)]
    if o.users_allowed:
        cols.append(Column("Active users who can run it", 14))
    cols.append(Column("Note", 50, wrap=True))
    if not o.ui:
        _table(ws, cols, [], NOT_SYNCED_UI)
        return
    granted = _grant_names(o, "FlowDefinition")
    rows = []
    for f in sorted(o.ui.get("flows", []), key=lambda f: (f.get("label") or f["api_name"]).lower()):
        self_starting = f.get("trigger_type") in ua.SELF_STARTING
        row = [f.get("label") or f["api_name"], f["api_name"], f.get("process_type") or "", f.get("trigger_type") or "",
               _yes(f.get("active")), "Not read" if f.get("restricted") is None else _yes(f.get("restricted")),
               ", ".join(sorted(granted.get(f["id"], [])))]
        if o.users_allowed:
            row.append(None if self_starting else runs.get(f["id"], 0))
        if self_starting:
            note = "Starts on its own; it doesn't depend on who can run it."
        elif f.get("restricted"):
            note = "Only through the explicit access listed."
        elif f.get("restricted") is None:
            note = "Whether it requires explicit access wasn't read."
        else:
            note = "Also anyone with the Run Flows permission or the Flow User setting."
        row.append(note)
        rows.append(row)
    _table(ws, cols, rows, "No flows.")


def _flow_rows(o: _Org) -> List[Tuple[Dict[str, Any], Optional[Dict[str, Any]]]]:
    details = (o.ui or {}).get("flow_details") or {}
    return [(f, details.get(f["id"])) for f in sorted((o.ui or {}).get("flows", []),
                                                       key=lambda f: (f.get("label") or f["api_name"]).lower())]


def _flow_elements(ws: Worksheet, o: _Org) -> None:
    """Element counts per flow (parity #9), from the active version."""
    flows = _flow_rows(o)
    used = [(k, label) for k, label in md_detail.FLOW_ELEMENTS if any((d or {}).get("elements", {}).get(k) for _, d in flows)]
    cols = [Column("Flow", 32), Column("API name", 30), Column("Start object", 18), Column("Runs when", 16)]
    cols += [Column(label, 10) for _, label in used]
    cols += [Column("Subflows called", 30, wrap=True), Column("Actions called", 40, wrap=True)]
    if not o.ui:
        _table(ws, cols, [], NOT_SYNCED_UI)
        return
    rows = []
    for f, d in flows:
        if d is None:
            rows.append([f.get("label") or f["api_name"], f["api_name"], "", "Not read (no active version, or not read)"]
                        + [None] * len(used) + ["", ""])
            continue
        rows.append([f.get("label") or f["api_name"], f["api_name"], d.get("start_object") or "",
                     {"Create": "Created", "Update": "Updated", "CreateAndUpdate": "Created or updated",
                      "Delete": "Deleted"}.get(d.get("record_trigger") or "", d.get("record_trigger") or f.get("trigger_type") or "")]
                    + [d["elements"].get(k, 0) for k, _ in used]
                    + [", ".join(d.get("subflows") or []),
                       ", ".join(f"{a['name']} ({a['type']})" if a.get("type") else a["name"] for a in d.get("actions") or [])])
    _table(ws, cols, rows, "No flows.")


def _flow_fields(ws: Worksheet, o: _Org) -> None:
    """The objects and fields each flow reads or writes (parity #9; feeds
    cleanup impact, item 13)."""
    cols = [Column("Flow", 32), Column("Element", 30), Column("Operation", 16), Column("Object", 22), Column("Field", 28),
            Column("Used as", 14), Column("Note", 44, wrap=True)]
    if not o.ui:
        _table(ws, cols, [], NOT_SYNCED_UI)
        return
    rows = []
    for f, d in _flow_rows(o):
        name = f.get("label") or f["api_name"]
        for op in (d or {}).get("operations", []):
            obj = op.get("object") or ""
            note = f"Whole record from the variable {op['from_variable']}" if op.get("from_variable") else ""
            written = op["operation"] in ("Create", "Update")
            uses = [(fld, "Written" if written else "Read") for fld in op.get("fields", [])]
            uses += [(fld, "Condition") for fld in op.get("filter_fields", [])]
            if not uses:
                rows.append([name, op.get("element") or "", op["operation"], obj, "", "", note])
            for fld, use in uses:
                rows.append([name, op.get("element") or "", op["operation"], obj, fld, use, note])
    _table(ws, cols, rows, "No record operations read from any flow.")


def _approval_steps(ws: Worksheet, o: _Org) -> None:
    cols = [Column("Approval process", 30), Column("Object", 20), Column("State", 10), Column("Step", 30),
            Column("API name", 28), Column("Description", 40, wrap=True), Column("Entry criteria and approvers", 44, wrap=True)]
    steps = _docs_list(o, "approval_steps")
    if steps is None:
        _table(ws, cols, [], NOT_SYNCED_DOCS)
        return
    procs = {p.get("api_name"): p for p in _docs_list(o, "approval_processes") or []}
    not_read = "Not read: they need the Metadata API, which isn't confirmed to work with Newton's permissions."
    rows = [[(procs.get(s.get("process")) or {}).get("name") or s.get("process"), (procs.get(s.get("process")) or {}).get("object"),
             (procs.get(s.get("process")) or {}).get("state"), s.get("name"), s.get("api_name"), s.get("description") or "",
             not_read] for s in steps]
    _table(ws, cols, rows, "No approval steps.")


def _cmdt_types(ws: Worksheet, o: _Org, enabled: Set[str]) -> None:
    cols = [Column("Type", 34), Column("Label", 30), Column("Records", 44, wrap=True)]
    types = _docs_list(o, "custom_metadata_types")
    if types is None:
        _table(ws, cols, [], NOT_SYNCED_DOCS)
        return
    stored = cmdt.keep_enabled((o.docs or {}).get("custom_metadata_records"), enabled)
    rows = []
    for t in sorted(types, key=lambda t: (t["name"] not in enabled, t["name"].lower())):
        if t["name"] not in enabled:
            what = "Not read: an admin can turn this type on in Privacy & Data."
        elif t["name"] in stored:
            n = len(stored[t["name"]]["records"])
            what = f"{n} {'record' if n == 1 else 'records'} (Custom Metadata Records sheet)"
        else:
            what = "Turned on; read from the next sync."
        rows.append([t["name"], t.get("label") or "", what])
    _table(ws, cols, rows, "No custom metadata types.")


def _cmdt_records(ws: Worksheet, o: _Org, enabled: Set[str]) -> None:
    """One row per record and field, for the types turned on now."""
    cols = [Column("Type", 30), Column("Record", 28), Column("Label", 28), Column("Namespace", 12), Column("Field", 28),
            Column("Value", 50, wrap=True)]
    stored = cmdt.keep_enabled((o.docs or {}).get("custom_metadata_records"), enabled)
    rows = []
    for t in sorted(stored, key=str.lower):
        for r in sorted(stored[t]["records"], key=lambda r: (r.get("name") or "").lower()):
            for f in stored[t]["fields"]:
                v = (r.get("values") or {}).get(f)
                rows.append([t, r.get("name"), r.get("label"), r.get("namespace") or "", f,
                             _yes(v) if isinstance(v, bool) else v])
    _table(ws, cols, rows, "No custom metadata records: no type is turned on, or none was read yet."
           if not enabled else "No records read yet for the types turned on: sync the org.")


def _docs_list(o: _Org, key: str) -> Optional[List[Dict[str, Any]]]:
    return None if o.docs is None else o.docs.get(key)


def _coverage(o: _Org) -> Dict[Tuple[str, str], Dict[str, Any]]:
    return {(c.get("type") or "", c.get("name") or ""): c for c in _docs_list(o, "coverage") or []}


def _pct(c: Optional[Dict[str, Any]]) -> Optional[float]:
    if not c or c.get("covered") is None or c.get("uncovered") is None:
        return None
    total = c["covered"] + c["uncovered"]
    return round(100 * c["covered"] / total, 1) if total else None


def _apex(ws: Worksheet, o: _Org, granted_users: Dict[str, int]) -> None:
    cols = [Column("Class", 32), Column("Namespace", 14), Column("API version", 11), Column("Status", 10),
            Column("Valid", 8), Column("Sharing", 18), Column("Size (characters)", 12), Column("Test class", 10),
            Column("Coverage %", 11, fmt='0.0'), Column("Lines covered", 11), Column("Lines not covered", 11),
            Column("Access granted by", 50, wrap=True)]
    if o.users_allowed:
        cols.append(Column("Active users granted it", 14))
    detail = _docs_list(o, "apex_classes")
    listed = detail if detail is not None else (o.ui or {}).get("apex_classes")
    if listed is None:
        _table(ws, cols, [], NOT_SYNCED_DOCS)
        return
    granted = _grant_names(o, "ApexClass")
    coverage = _coverage(o)
    rows = []
    for c in sorted(listed, key=lambda c: (c.get("name") or "").lower()):
        cov = coverage.get(("class", c.get("name") or ""))
        row = [c.get("name"), c.get("namespace") or "", c.get("api_version"), c.get("status") or "",
               _yes(c.get("valid")), c.get("sharing") or "", c.get("length"),
               _yes(cov.get("is_test")) if cov else "", _pct(cov),
               cov.get("covered") if cov else None, cov.get("uncovered") if cov else None,
               ", ".join(sorted(granted.get(c.get("id") or "", []))) or ""]
        if o.users_allowed:
            row.append(granted_users.get(c.get("id") or "", 0))
        rows.append(row)
    _table(ws, cols, rows, "No Apex classes.")


def _apex_triggers(ws: Worksheet, o: _Org) -> None:
    cols = [Column("Trigger", 30), Column("Object", 24), Column("Events", 50, wrap=True), Column("Status", 10),
            Column("Valid", 8), Column("API version", 11), Column("Coverage %", 11, fmt='0.0'),
            Column("Lines covered", 11), Column("Lines not covered", 11)]
    triggers = _docs_list(o, "apex_triggers")
    if triggers is None:
        _table(ws, cols, [], NOT_SYNCED_DOCS)
        return
    coverage = _coverage(o)
    rows = []
    for t in sorted(triggers, key=lambda t: ((t.get("object") or ""), (t.get("name") or "").lower())):
        cov = coverage.get(("trigger", t.get("name") or ""))
        rows.append([t.get("name"), t.get("object") or "", ", ".join(t.get("events") or []), t.get("status") or "",
                     _yes(t.get("valid")), t.get("api_version"), _pct(cov),
                     cov.get("covered") if cov else None, cov.get("uncovered") if cov else None])
    _table(ws, cols, rows, "No Apex triggers.")


_PERMISSION_REF = re.compile(r"\$Permission\.([A-Za-z0-9_]+)")


def _validation_rules(ws: Worksheet, o: _Org, permission_users: Dict[str, int]) -> None:
    cols = [Column("Object", 22), Column("Rule", 32), Column("Active", 8), Column("Formula (error when true)", 70, wrap=True),
            Column("Error message", 44, wrap=True), Column("Shown on", 20),
            Column("Custom permissions it checks", 40, wrap=True), Column("Description", 36, wrap=True)]
    rules = _docs_list(o, "validation_rules")
    if rules is None:
        _table(ws, cols, [], NOT_SYNCED_DOCS)
        return
    perm_ids = {p.get("name"): p.get("id") for p in _docs_list(o, "custom_permissions") or []}

    def checks(formula: Optional[str]) -> str:
        out = []
        for name in sorted(set(_PERMISSION_REF.findall(formula or ""))):
            pid = perm_ids.get(name)
            if o.users_allowed and pid:
                n = permission_users.get(pid, 0)
                out.append(f"{name} ({n} active user{'s' if n != 1 else ''} hold it)")
            else:
                out.append(name)
        return "; ".join(out)

    rows = [[r.get("object") or "", r.get("name"), _yes(r.get("active")),
             r.get("formula") if r.get("formula") is not None else "Not read",
             r.get("error_message") or "", r.get("error_field") or "Top of page", checks(r.get("formula")),
             r.get("description") or ""]
            for r in sorted(rules, key=lambda r: ((r.get("object") or ""), (r.get("name") or "").lower()))]
    _table(ws, cols, rows, "No validation rules.")


def _rt_picklists(ws: Worksheet, o: _Org) -> None:
    cols = [Column("Object", 22), Column("Record type", 26), Column("Field", 26), Column("Values available", 90, wrap=True),
            Column("Values left out", 50, wrap=True)]
    data = None if o.docs is None else o.docs.get("record_type_picklists")
    if data is None:
        _table(ws, cols, [], NOT_SYNCED_DOCS)
        return
    described = _described_objects(o)
    rows = []
    for obj in sorted(data):
        for rt in sorted(data[obj]):
            for fld, values in sorted(data[obj][rt].items()):
                every = [v.get("value") for v in ((described.get(obj) or {}).get("fields", {}).get(fld) or {}).get("picklistValues") or []]
                left_out = [v for v in every if v not in values]
                rows.append([obj, rt, fld, ", ".join(values), ", ".join(left_out)])
    _table(ws, cols, rows, "No record-type-specific picklist values.")


def _tabs(ws: Worksheet, o: _Org) -> None:
    settings = _docs_list(o, "tab_settings")
    if settings is None:
        _table(ws, [Column("Tab", 30)], [], NOT_SYNCED_DOCS)
        return
    columns = [(h, src.components[0].ps_id) for kind, h, src in o.columns() if kind != "permission_set_group"]
    by: Dict[str, Dict[str, str]] = defaultdict(dict)
    for t in settings:
        by[t.get("tab") or ""][t.get("parent_id") or ""] = {"DefaultOn": "Default On", "DefaultOff": "Default Off"}.get(
            t.get("visibility") or "", t.get("visibility") or "")
    used = [(h, pid) for h, pid in columns if any(pid in v for v in by.values())]
    head = [Column("Tab", 30)] + [Column(h, 14) for h, _ in used]
    rows = [[tab] + [by[tab].get(pid, "") for _, pid in used] for tab in sorted(by)]
    _table(ws, head, rows, "No tab settings.")
    for i in range(2, len(used) + 2):
        ws.cell(row=1, column=i).alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"


def _login_restrictions(ws: Worksheet, o: _Org) -> None:
    cols = [Column("Profile", 28), Column("Login IP ranges", 60, wrap=True), Column("Login hours", 60, wrap=True)]
    if o.users_allowed:
        cols.append(Column("Active users with the profile", 14))
    if not o.ui:
        _table(ws, cols, [], NOT_SYNCED_UI)
        return
    logins = o.ui.get("profile_logins") or {}
    per_profile = _active_by_profile(o)

    def hhmm(m: Any) -> str:
        try:
            return f"{int(m) // 60:02d}:{int(m) % 60:02d}"
        except (TypeError, ValueError):
            return str(m)

    rows = []
    for pid, name in sorted(o.profiles.items(), key=lambda p: p[1].lower()):
        lg = logins.get(pid) or {}
        ranges = "; ".join(f"{r.get('start')} to {r.get('end')}" + (f" ({r['description']})" if r.get("description") else "")
                           for r in lg.get("ip_ranges") or [])
        hours = "; ".join(f"{DAY_ABBR.get(d, d)} {hhmm(h.get('start'))} to {hhmm(h.get('end'))}"
                          for d, h in (lg.get("hours") or {}).items())
        row: List[Any] = [name, ranges or "No restriction", hours or "No restriction"]
        if o.users_allowed:
            row.append(per_profile.get(pid, 0))
        rows.append(row)
    _table(ws, cols, rows, "No profiles synced.")


def _simple(ws: Worksheet, o: _Org, key: str, cols: List[Column], row_of, sort_key, empty: str) -> None:
    items = _docs_list(o, key)
    if items is None:
        _table(ws, cols, [], NOT_SYNCED_DOCS)
        return
    _table(ws, cols, [row_of(x) for x in sorted(items, key=sort_key)], empty)


def _granted_sheet(ws: Worksheet, o: _Org, key: str, kind: str, cols: List[Column], row_of, users: Dict[str, int],
                   empty: str) -> None:
    """A list whose items are granted through SetupEntityAccess, with who."""
    granted = _grant_names(o, kind)
    full = cols + [Column("Access granted by", 50, wrap=True)] + ([Column("Active users granted it", 14)] if o.users_allowed else [])

    def row(x: Dict[str, Any]) -> List[Any]:
        r = row_of(x)
        if kind != "ApexPage" or x.get("kind") == "Page":
            r.append(", ".join(sorted(granted.get(x.get("id") or "", []))))
            if o.users_allowed:
                r.append(users.get(x.get("id") or "", 0))
        else:
            r.append("Components aren't granted on their own")
            if o.users_allowed:
                r.append(None)
        return r

    _simple(ws, o, key, full, row, lambda x: ((x.get("kind") or ""), (x.get("name") or "").lower()), empty)


async def build_documentation_workbook(db: AsyncSession, org: Organization, now: Optional[datetime] = None) -> bytes:
    now = now or datetime.now(timezone.utc)
    o = await _load(db, org)
    wb = Workbook()
    sheets: Dict[str, Worksheet] = {}
    for i, title in enumerate(SHEETS):
        if i == 0:
            ws = wb.active
            assert ws is not None
            ws.title = title
        else:
            ws = wb.create_sheet(title)
        sheets[title] = ws
    flow_runs, granted_users = _per_user_runs(o)
    _about(sheets["About"], org, o, now)
    _objects(sheets["Objects"], o)
    _fields(sheets["Fields"], o)
    _picklists(sheets["Picklist Values"], o)
    _matrix(sheets["Object Access"], o, "objects")
    _matrix(sheets["Field Access"], o, "fields")
    _assignments(sheets["Assignments"], o)
    _profile_users(sheets["Profile Users"], o)
    _record_types(sheets["Record Types"], o)
    _layouts(sheets["Layouts"], o)
    _layout_fields(sheets["Fields on Layouts"], o)
    _unused_layouts(sheets["Unused Layouts"], o)
    _record_pages(sheets["Record Pages"], o)
    _rt_picklists(sheets["Picklists by Record Type"], o)
    _validation_rules(sheets["Validation Rules"], o, granted_users["CustomPermission"])
    _tabs(sheets["Tab Visibility"], o)
    _login_restrictions(sheets["Login Restrictions"], o)
    _simple(sheets["Buttons and Links"], o, "web_links",
            [Column("Object", 22), Column("Label", 28), Column("API name", 28), Column("Shown as", 18),
             Column("Link type", 12), Column("Opens", 14)],
            lambda w: [w.get("object"), w.get("label"), w.get("name"), w.get("display_type"), w.get("link_type"),
                       w.get("open_type")],
            lambda w: ((w.get("object") or ""), (w.get("label") or "").lower()), "No custom buttons or links.")
    _simple(sheets["List Views"], o, "list_views", [Column("Object", 22), Column("List view", 32), Column("API name", 32)],
            lambda v: [v.get("object"), v.get("name"), v.get("api_name")],
            lambda v: ((v.get("object") or ""), (v.get("name") or "").lower()), "No list views.")
    _simple(sheets["Field Sets"], o, "field_sets", [Column("Field set", 32), Column("Label", 32), Column("Description", 60, wrap=True)],
            lambda f: [f.get("name"), f.get("label"), f.get("description") or ""],
            lambda f: (f.get("name") or "").lower(), "No field sets.")
    _granted_sheet(sheets["Apps"], o, "apps", "TabSet",
                   [Column("App", 28), Column("API name", 28), Column("Namespace", 14), Column("Description", 40, wrap=True)],
                   lambda a: [a.get("label"), a.get("name"), a.get("namespace") or "", a.get("description") or ""],
                   granted_users["TabSet"], "No apps.")
    _flows(sheets["Flows"], o, flow_runs)
    _flow_elements(sheets["Flow Elements"], o)
    _flow_fields(sheets["Flow Fields"], o)
    _approval_steps(sheets["Approval Steps"], o)
    _simple(sheets["Approval Processes"], o, "approval_processes",
            [Column("Approval process", 32), Column("API name", 30), Column("Object", 22), Column("State", 10),
             Column("Record lock", 12), Column("Description", 50, wrap=True)],
            lambda p: [p.get("name"), p.get("api_name"), p.get("object"), p.get("state"), p.get("lock_type"),
                       p.get("description") or ""],
            lambda p: ((p.get("object") or ""), (p.get("name") or "").lower()), "No approval processes.")
    _apex(sheets["Apex Classes"], o, granted_users["ApexClass"])
    _apex_triggers(sheets["Apex Triggers"], o)
    _granted_sheet(sheets["Visualforce"], o, "visualforce", "ApexPage",
                   [Column("Type", 11), Column("Name", 30), Column("Label", 30), Column("Namespace", 14),
                    Column("API version", 11), Column("Description", 40, wrap=True)],
                   lambda v: [v.get("kind"), v.get("name"), v.get("label"), v.get("namespace") or "", v.get("api_version"),
                              v.get("description") or ""],
                   granted_users["ApexPage"], "No Visualforce pages or components.")
    _simple(sheets["Lightning Components"], o, "lightning_components",
            [Column("Type", 24), Column("Name", 28), Column("Label", 28), Column("Namespace", 14), Column("API version", 11),
             Column("Exposed in App Builder", 12), Column("Description", 44, wrap=True)],
            lambda c: [c.get("kind"), c.get("name"), c.get("label"), c.get("namespace") or "", c.get("api_version"),
                       _yes(c.get("exposed")), c.get("description") or ""],
            lambda c: ((c.get("kind") or ""), (c.get("name") or "").lower()), "No Lightning components.")
    _simple(sheets["Static Resources"], o, "static_resources",
            [Column("Name", 28), Column("Namespace", 14), Column("Content type", 24), Column("Size (bytes)", 14, fmt="#,##0"),
             Column("Cache control", 12), Column("Description", 44, wrap=True)],
            lambda r: [r.get("name"), r.get("namespace") or "", r.get("content_type"), r.get("size"), r.get("cache_control"),
                       r.get("description") or ""],
            lambda r: (r.get("name") or "").lower(), "No static resources.")
    _granted_sheet(sheets["Custom Permissions"], o, "custom_permissions", "CustomPermission",
                   [Column("Custom permission", 30), Column("Label", 30), Column("Namespace", 14),
                    Column("Description", 44, wrap=True)],
                   lambda p: [p.get("name"), p.get("label"), p.get("namespace") or "", p.get("description") or ""],
                   granted_users["CustomPermission"], "No custom permissions.")
    _granted_sheet(sheets["External Data Sources"], o, "external_data_sources", "ExternalDataSource",
                   [Column("External data source", 30), Column("Label", 30), Column("Type", 18), Column("Namespace", 14)],
                   lambda x: [x.get("name"), x.get("label"), x.get("type"), x.get("namespace") or ""],
                   granted_users["ExternalDataSource"], "No external data sources.")
    on = cmdt.enabled_types(org)
    _cmdt_types(sheets["Custom Metadata Types"], o, on)
    _cmdt_records(sheets["Custom Metadata Records"], o, on)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

