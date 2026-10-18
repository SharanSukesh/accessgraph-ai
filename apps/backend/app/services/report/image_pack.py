"""Image pack (roadmap item 3): the Health Report's visuals as ready-to-drop
images, one at a time or as a set, in a ZIP or a PowerPoint file.

Each visual is an HTML page drawn with the PDF report's charts and
styles, rendered by WeasyPrint and rasterised with pypdfium2:

  slide  16:9, laid out at 800x450 and rendered at 3840x2160, with a
         title and a short caption; no footer unless the org name is
         asked for (org_name).
  crop   the visual alone (no slide frame), for pasting into an existing
         slide: laid out 680 wide, rendered at the same scale and cropped
         to its content.

The images go into the consultant's own branded slides, so they carry no
logos and no "Demo data" label (demo files keep the demo- name prefix and
the ZIP's DEMO-DATA.txt; the app shows its own demo badge). The Summary
visual puts the key figures on one slide. PNG can have a transparent
background so the image sits on the client's own template; JPEG is always
on white. Names are whatever Newton stored, so a masked org's images show
aliases; user-level visuals (and the Summary's user-level tiles) are left
out when the org's privacy level excludes user records. A visual whose
module hasn't run is skipped with a note. Nothing here calls Salesforce.
"""
from __future__ import annotations

import io
import zipfile
from collections import Counter
from dataclasses import dataclass, field
from datetime import datetime, timezone
from typing import Any, Callable, Dict, List, Optional, Tuple

from app.services.report import charts
from app.services.report.context import RISK_LEVELS, ReportContext

SLIDE_W, SLIDE_H = 800, 450          # CSS px (1/96 in)
CROP_W = 680
PX_PER_PT = 96 / 72
SCALE = 3840 / SLIDE_W                # image px per CSS px: 3840x2160 slides
CROP_PAD = 12                         # CSS px of margin kept around a crop
FORMATS = ("png", "jpeg")
LAYOUTS = ("slide", "crop")

USERS_HIDDEN = "Left out: this client's privacy level excludes user records."
SEVERITIES = ("critical", "high", "medium", "low", "info")


@dataclass
class Visual:
    key: str
    title: str
    user_level: bool = False


VISUALS: List[Visual] = [
    Visual("summary", "Summary"),
    Visual("health_score", "Org health score"),
    Visual("savings", "Estimated annual savings"),
    Visual("risk_distribution", "User access risk", user_level=True),
    Visual("top_anomalies", "Top access anomalies", user_level=True),
    Visual("license_fit", "Licence fit", user_level=True),
    Visual("compliance", "Compliance control indicators"),
    Visual("critical_objects", "Critical objects"),
    Visual("code_coverage", "Code coverage"),
    Visual("api_usage", "API usage"),
]
VISUAL_BY_KEY = {v.key: v for v in VISUALS}


@dataclass
class PackSource:
    """Everything the visuals draw from; built by gather_pack_source()."""
    org_name: str
    is_demo: bool = False
    allow_users: bool = True
    snapshot: Any = None
    findings: List[Any] = field(default_factory=list)
    ctx: ReportContext = field(default_factory=ReportContext)
    critical_objects: List[Dict[str, Any]] = field(default_factory=list)
    data_quality_run: bool = False

    @property
    def metrics(self) -> Dict[str, Any]:
        return (getattr(self.snapshot, "metrics", None) or {}) if self.snapshot is not None else {}


@dataclass
class ImageOptions:
    fmt: str = "png"
    transparent: bool = False
    layout: str = "slide"
    # Slides carry no footer unless asked: then the org's name, bottom left.
    org_name: bool = False

    def __post_init__(self) -> None:
        self.fmt = "jpeg" if self.fmt in ("jpg", "jpeg") else self.fmt
        if self.fmt not in FORMATS:
            raise ValueError(f"format must be one of {FORMATS}")
        if self.layout not in LAYOUTS:
            raise ValueError(f"layout must be one of {LAYOUTS}")
        if self.fmt == "jpeg":
            self.transparent = False  # JPEG has no alpha channel


# ---------------------------------------------------------------- helpers


def _esc(v: object) -> str:
    import html
    return html.escape(str(v), quote=True)


def _money(cents: int) -> str:
    return f"${cents / 100:,.0f}"


def _compact(n: float) -> str:
    n = float(n or 0)
    for div, unit in ((1e6, "M"), (1e3, "K")):
        if abs(n) >= div:
            return f"{n / div:.1f}{unit}".replace(".0" + unit, unit)
    return f"{n:,.0f}"


def _band(score: float) -> str:
    from app.services.org_analyzer import _health_band
    return _health_band(int(score)).capitalize()


def _value(e: Any) -> str:
    return str(getattr(e, "value", e) or "").lower()


# ---------------------------------------------------------------- visuals
# Each returns (caption, body html) or a skip note (str).

Body = Tuple[str, str]


def _tile(label: str, value: str, note: str = "", wide: bool = False) -> str:
    return (f'<div class="tile{" wide" if wide else ""}"><div class="tile-label">{_esc(label)}</div>'
            f'{value}{note}</div>')


def _tile_missing(label: str, text: str) -> str:
    return _tile(label, f'<p class="tile-missing">{_esc(text)}</p>')


def _severity_strip(counts: Dict[str, int]) -> str:
    total = sum(counts.values())
    segs = "".join(
        f'<span style="width:{counts[s] / total * 100:.2f}%;background:{charts.SEVERITY_RAMP[s]}"></span>'
        for s in SEVERITIES if counts.get(s)
    )
    keys = "".join(
        f'<span class="key"><i style="background:{charts.SEVERITY_RAMP[s]}"></i>{counts[s]} {s.capitalize()}</span>'
        for s in SEVERITIES if counts.get(s)
    )
    return f'<div class="strip">{segs}</div><div class="keys">{keys}</div>'


def _summary(src: PackSource):
    """The key figures on one slide: one tile per module, and a module that
    hasn't run (or that the privacy level hides) says so in its tile."""
    m = src.metrics
    if src.snapshot is None or m.get("org_health_score") is None:
        return "Run the Health Report first."
    tiles: List[str] = []

    score = float(m["org_health_score"])
    color = charts.SERIES if score >= 75 else (charts.WARNING if score >= 60 else charts.CRITICAL)
    bar = (f'<svg width="200" height="7" viewBox="0 0 200 7"><rect width="200" height="7" rx="3.5" fill="{charts.TRACK}"/>'
           f'<rect width="{max(200 * score / 100, 7):.1f}" height="7" rx="3.5" fill="{color}"/></svg>')
    tiles.append(_tile("Health score", f'<div class="tile-value">{score:.0f}<span class="of">/100</span></div>',
                       f'<p class="tile-note">{_esc(_band(score))}</p>{bar}'))

    sev = Counter(_value(f.severity) for f in src.findings)
    tiles.append(_tile("Findings", f'<div class="tile-value">{len(src.findings)}</div>',
                       _severity_strip(sev) if src.findings else '<p class="tile-note">None open</p>'))

    savings = sum(f.estimated_annual_savings_cents or 0 for f in src.findings)
    lf = src.ctx.license_fit
    if not src.allow_users:
        lf_note = "License Fit left out (privacy level)"
    elif lf is None or not lf.users_assessed:
        lf_note = "License Fit not run yet"
    else:
        lf_note = f"License Fit: {_money(lf.total_annual_savings_cents or 0)} a year from right-sizing"
    tiles.append(_tile("Estimated savings", f'<div class="tile-value">{_money(savings)}<span class="unit"> a year</span></div>',
                       f'<p class="tile-note">{_esc(lf_note)}</p>'))

    if not src.allow_users:
        tiles.append(_tile_missing("High-risk users", USERS_HIDDEN))
        tiles.append(_tile_missing("Flagged anomalies", USERS_HIDDEN))
    else:
        counts = src.ctx.risk_counts
        if counts:
            high = counts.get("critical", 0) + counts.get("high", 0)
            tiles.append(_tile("High-risk users", f'<div class="tile-value">{high:,}</div>',
                               f'<p class="tile-note">{counts.get("critical", 0):,} critical · {counts.get("high", 0):,} high, '
                               f'of {sum(counts.values()):,} scored</p>'))
        else:
            tiles.append(_tile_missing("High-risk users", "No risk scores yet: sync the org."))
        cats = src.ctx.anomaly_by_category
        flagged = sum(cats.values())
        if flagged:
            top = "".join(f'<li><b>#{a["rank"] or "?"}</b> {_esc(a["name"])} <span>{_esc(a["reason"])}</span></li>'
                          for a in src.ctx.flagged_access[:3])
            split = f'{cats.get("access", 0):,} access · {cats.get("session", 0):,} session'
            tiles.append(_tile("Flagged anomalies", f'<div class="tile-value">{flagged:,}<span class="unit"> {split}</span></div>',
                               f'<ol class="top">{top}</ol>' if top else ""))
        else:
            tiles.append(_tile_missing("Flagged anomalies", "None flagged: sync the org to run the analysis."))

    rows = src.ctx.compliance
    if rows:
        lines = "".join(
            f'<li>{_esc(c["framework"])} <b>{c["passed"]} of {c["passed"] + c["failed"]}</b></li>' for c in rows[:4]
        )
        passed = sum(c["passed"] for c in rows)
        judged = sum(c["passed"] + c["failed"] for c in rows)
        tiles.append(_tile("Compliance controls passing",
                           f'<div class="tile-value">{passed}<span class="of"> of {judged}</span></div>',
                           f'<ul class="list">{lines}</ul>'))
    else:
        tiles.append(_tile_missing("Compliance controls passing", "No compliance scorecard run yet."))

    cov = m.get("code_coverage") or {}
    if cov.get("org_wide_pct") is not None:
        gate = float(cov["gate_pct"])
        verdict = f"Below the {gate:g}% deployment gate" if cov.get("below_gate") else f"Meets the {gate:g}% deployment gate"
        tiles.append(_tile("Code coverage", f'<div class="tile-value">{float(cov["org_wide_pct"]):g}%</div>',
                           f'<p class="tile-note">{_esc(verdict)}</p>'))
    else:
        tiles.append(_tile_missing("Code coverage", "No coverage data in the latest Health Report."))

    u = m.get("api_usage") or {}
    peak, limit = u.get("peak") or {}, (u.get("limit") or {}).get("max")
    if u.get("status") == "ok" and peak and limit:
        days = u.get("days_collected") or 0
        span = f"busiest of {days} days" if days > 1 else "the only day collected"
        tiles.append(_tile("API peak day",
                           f'<div class="tile-value">{peak["pct_of_limit"]:g}%<span class="unit"> of the daily limit</span></div>',
                           f'<p class="tile-note">{_esc(peak["date"])}, {span}: {peak["counting"]:,} of {limit:,} requests</p>'))
    else:
        tiles.append(_tile_missing("API peak day", u.get("reason") or "No API usage in the latest Health Report."))

    if not src.data_quality_run:
        tiles.append(_tile_missing("Critical objects", "Data Quality not run yet."))
    else:
        objs = src.critical_objects
        order = {"stale": 0, "mostly_empty": 1, "never_populated": 2}
        objs = sorted(objs, key=lambda o: (min(order.get(f["code"], 9) for f in o["flags"]), -(o["record_count"] or 0)))
        names = ", ".join(o["label"] or o["object"] for o in objs[:3]) + (f" and {len(objs) - 3} more" if len(objs) > 3 else "")
        tiles.append(_tile("Critical objects", f'<div class="tile-value">{len(objs)}</div>',
                           f'<p class="tile-note">{_esc(names) if objs else "None in the latest Data Quality run"}</p>'))

    return ("The latest Health Report and each module's latest run, on one slide.",
            f'<div class="grid">{"".join(tiles)}</div>')



def _health_score(src: PackSource):
    m = src.metrics
    if src.snapshot is None or m.get("org_health_score") is None:
        return "Run the Health Report first."
    score = float(m["org_health_score"])
    color = charts.SERIES if score >= 75 else (charts.WARNING if score >= 60 else charts.CRITICAL)
    fams = (m.get("org_health_rubric") or {}).get("families") or []
    bars = charts.horizontal_bars(
        [(f["label"], f["deduction"]) for f in fams if f.get("deduction")],
        width=420, label_width=170, value_label=lambda v: f"−{v:.0f}",
    )
    body = f"""
      <div class="row">
        <div class="hero">
          <div class="hero-value">{score:.0f}<span class="of">/100</span></div>
          <div class="hero-band">{_esc(_band(score))}</div>
          <svg width="230" height="10" viewBox="0 0 230 10"><rect width="230" height="10" rx="5" fill="{charts.TRACK}"/>
          <rect width="{max(230 * score / 100, 10):.1f}" height="10" rx="5" fill="{color}"/></svg>
          <p class="note">{len(src.findings)} findings</p>
        </div>
        {f'<div class="grow"><p class="label">Points lost by family (each capped at 20)</p>{bars}</div>' if bars else ''}
      </div>"""
    return "100 minus severity-weighted deductions for each finding; one family can lower the score by 20 at most.", body


def _savings(src: PackSource):
    if src.snapshot is None:
        return "Run the Health Report first."
    rows = sorted((f for f in src.findings if (f.estimated_annual_savings_cents or 0) > 0),
                  key=lambda f: -(f.estimated_annual_savings_cents or 0))
    if not rows:
        return "No savings estimated in the latest Health Report."
    total = sum(f.estimated_annual_savings_cents or 0 for f in rows)
    items = [(f.title if len(f.title) <= 56 else f.title[:54] + "…", f.estimated_annual_savings_cents / 100) for f in rows[:7]]
    rest = sum(f.estimated_annual_savings_cents for f in rows[7:])
    if rest:
        items.append((f"{len(rows) - 7} more findings", rest / 100))
    bars = charts.horizontal_bars(items, width=720, label_width=340, value_label=lambda v: _money(int(v * 100)))
    body = f'<p class="big">{_money(total)}<span class="unit"> a year</span></p>{bars}'
    return "Published list prices; negotiated rates replace them. Ignored findings are left out.", body


def _risk_distribution(src: PackSource):
    counts = src.ctx.risk_counts
    if not counts:
        return "No risk scores yet: sync the org to run the analysis."
    total = sum(counts.values())
    color = {"critical": "critical", "high": "high", "medium": "medium", "low": "info"}
    bar = charts.stacked_bar(
        [(lvl.capitalize(), counts.get(lvl, 0), charts.SEVERITY_RAMP[color[lvl]]) for lvl in RISK_LEVELS], width=640,
    )
    tiles = "".join(
        f'<div class="stat"><div class="stat-label">{lvl.capitalize()}</div><div class="stat-value">{counts.get(lvl, 0):,}</div></div>'
        for lvl in RISK_LEVELS
    )
    return f"{total:,} users scored on what their access lets them do.", f'<div class="tiles">{tiles}</div>{bar}'


def _top_anomalies(src: PackSource):
    rows = src.ctx.flagged_access
    if not rows:
        return "No access anomalies flagged: sync the org to run the analysis."
    scored = rows[0].get("users_scored")
    body = "".join(
        f"<tr><td class='num'>#{a['rank'] or '?'}</td><td>{_esc(a['name'])}</td><td>{_esc(a['reason'])}</td>"
        f"<td>{_esc((a['base_risk_level'] or 'not scored').capitalize())}</td></tr>"
        for a in rows[:8]
    )
    caption = (f"The {len(rows)} users whose access stands out most from their peers"
               + (f", out of {scored:,} scored" if scored else "") + ", in rank order. Base risk excludes the anomaly flag.")
    return caption, (f'<table class="data"><thead><tr><th class="num">Rank</th><th>User</th><th>Strongest reason</th>'
                     f'<th>Base risk</th></tr></thead><tbody>{body}</tbody></table>')


def _license_fit(src: PackSource):
    lf = src.ctx.license_fit
    if lf is None or not lf.users_assessed:
        return "Run License Fit first."
    bar = charts.horizontal_bars([
        ("Right-sized", lf.users_right_sized), ("Over-licensed", lf.users_overbuilt),
        ("Wrong cloud", lf.users_wrong_cloud), ("Under-used", lf.users_underused),
        ("Inactive but billed", lf.users_inactive_billed),
    ], width=640, label_width=200)
    body = f'<p class="big">{_money(lf.total_annual_savings_cents or 0)}<span class="unit"> a year from right-sizing</span></p>{bar}'
    return f"{lf.users_assessed:,} users assessed against the licence their work needs.", body


def _compliance(src: PackSource):
    rows = src.ctx.compliance
    if not rows:
        return "Run a compliance scorecard first."
    body = "".join(
        f"<tr><td>{_esc(c['framework'])}</td><td>{charts.meter((c['score_pct'] or 0) / 100, width=220)}</td>"
        f"<td class='num'>{(c['score_pct'] or 0):.0f}%</td>"
        f"<td class='num'>{c['passed']} passed · {c['failed']} failed · {c['na']} n/a</td></tr>"
        for c in rows
    )
    return ("Automated indicators based on org configuration. Not an audit opinion, attestation or certification.",
            f'<table class="data"><thead><tr><th>Framework</th><th>Controls passed</th><th class="num">Score</th>'
            f'<th class="num">Controls</th></tr></thead><tbody>{body}</tbody></table>')


def _critical_objects(src: PackSource):
    if not src.data_quality_run:
        return "Run Data Quality first."
    objs = src.critical_objects
    if not objs:
        return "No critical objects in the latest Data Quality run."
    order = {"stale": 0, "mostly_empty": 1, "never_populated": 2}
    objs = sorted(objs, key=lambda o: (min(order.get(f["code"], 9) for f in o["flags"]), -(o["record_count"] or 0)))
    body = "".join(
        f"<tr><td>{_esc(o['label'] or o['object'])}</td><td class='num'>{(o['record_count'] or 0):,}</td>"
        f"<td>{_esc(', '.join(f['label'] for f in o['flags']))}</td>"
        f"<td class='reason'>{_esc(o['flags'][0]['reason'])}</td></tr>"
        for o in objs[:7]
    )
    more = f'<p class="note">and {len(objs) - 7} more</p>' if len(objs) > 7 else ""
    caveat = " Counts marked as estimates or partial in Newton keep those caveats." if any(o.get("caveats") for o in objs) else ""
    return (f"Objects whose data is stale, mostly empty or has fields nobody fills in.{caveat}",
            f'<table class="data"><thead><tr><th>Object</th><th class="num">Records</th><th>Flags</th><th>Main reason</th>'
            f'</tr></thead><tbody>{body}</tbody></table>{more}')


def _code_coverage(src: PackSource):
    cov = src.metrics.get("code_coverage")
    if not cov or cov.get("org_wide_pct") is None:
        return "No coverage data in the latest Health Report."
    pct, gate = float(cov["org_wide_pct"]), float(cov["gate_pct"])
    w = 560
    color = charts.CRITICAL if cov.get("below_gate") else charts.SERIES
    verdict = f"Below the {gate:g}% deployment gate" if cov.get("below_gate") else f"Meets the {gate:g}% deployment gate"
    bar = (f'<svg width="{w}" height="40" viewBox="0 0 {w} 40"><rect y="8" width="{w}" height="12" rx="6" fill="{charts.TRACK}"/>'
           f'<rect y="8" width="{max(w * pct / 100, 12):.1f}" height="12" rx="6" fill="{color}"/>'
           f'<line x1="{w * gate / 100:.1f}" x2="{w * gate / 100:.1f}" y1="2" y2="26" stroke="{charts.INK}" stroke-width="2"/>'
           f'<text x="{w * gate / 100:.1f}" y="38" text-anchor="middle" font-size="10" fill="{charts.INK_2}" '
           f'style="{charts.FONT}">{gate:g}% gate</text></svg>')
    zero = len(cov.get("zero_classes") or []) + len(cov.get("zero_triggers") or [])
    days = cov.get("days_since_run")
    tiles = "".join([
        f'<div class="stat"><div class="stat-label">Below {gate:g}%</div><div class="stat-value">{cov.get("below_gate_items", 0)}'
        f'<span class="of"> of {cov.get("measured", 0)}</span></div></div>',
        f'<div class="stat"><div class="stat-label">At 0%</div><div class="stat-value">{zero}</div></div>',
        f'<div class="stat"><div class="stat-label">Last test run</div><div class="stat-value">'
        f'{"—" if days is None else f"{days} days ago"}</div></div>',
        f'<div class="stat"><div class="stat-label">Changed since</div><div class="stat-value">{cov.get("changed_since_run_count", 0)}</div></div>',
    ])
    body = f'<p class="big">{pct:g}%<span class="unit"> org-wide · {_esc(verdict)}</span></p>{bar}<div class="tiles">{tiles}</div>'
    return "Read from Salesforce's stored test results; Newton never runs tests.", body


def _api_usage(src: PackSource):
    """The Health Report's API usage card: headline, four figures, the daily
    chart against the limit and the top 3 sources."""
    u = src.metrics.get("api_usage")
    if not u or u.get("status") != "ok":
        return (u or {}).get("reason") or "No API usage in the latest Health Report."
    days = u.get("days_collected") or 0
    limit = u.get("limit") or {}
    peak = u.get("peak") or {}
    top = [a for a in u.get("by_app") or [] if not a.get("other")][:3]
    period = f"{u['first_day']} to {u['last_day']}" if days > 1 else f"{u['last_day']} only"
    head = f"{top[0]['name']} made {top[0]['share_pct']:g}% of requests." if top else ""

    def big(n: float) -> str:  # as the card: 1,015,000 -> 1.02M
        return f"{round(n / 10_000) / 100:g}M" if n >= 1e6 else _compact(n)

    def stat(label: str, value: str, note: str = "") -> str:
        return (f'<div class="stat"><div class="stat-label">{_esc(label)}</div><div class="stat-value">{value}</div>'
                + (f'<div class="stat-label">{_esc(note)}</div>' if note else "") + "</div>")
    today = (stat("Today so far", f'{big(limit["used"])}<span class="of"> of {big(limit["max"])}</span>',
                  f'{limit["used"] / limit["max"] * 100:.0f}% of the daily limit') if limit.get("max") else
             stat("Today so far", "—", "Limit not read"))
    peak_tile = (stat("Peak day", f'{peak["pct_of_limit"]:g}%<span class="of"> of limit</span>', peak["date"])
                 if peak.get("pct_of_limit") is not None else stat("Peak day", _compact(peak.get("counting") or 0), peak.get("date") or ""))
    tiles = "".join([
        stat(f"{days}-day total" if days > 1 else "Requests", big(u["total_requests"]), period),
        stat("Daily average", big(u.get("daily_average") or 0)),
        today,
        peak_tile,
    ])
    trend = u.get("trend") or []
    chart = charts.daily_bars([(t["date"][5:], t["counting"]) for t in trend], limit=limit.get("max"), width=720,
                              height=96, near_at=0.9, value_label=big) if len(trend) > 1 else ""
    sources = "".join(
        f"<tr><td>{_esc(a['name'])}</td><td>{charts.meter(a['share_pct'] / 100, width=160)}</td>"
        f"<td class='num'>{a['share_pct']:g}%</td><td class='num'>{big(a['requests'])}</td></tr>"
        for a in top
    )
    body = (f'<div class="tiles">{tiles}</div>'
            + (f'<p class="label">Requests counted against the limit, per day</p>{chart}' if chart else "")
            + (f'<table class="data compact"><thead><tr><th>Top sources</th><th></th><th class="num">Share</th>'
               f'<th class="num">Requests</th></tr></thead><tbody>{sources}</tbody></table>' if sources else "")
            + f'<p class="note">{_esc(u.get("coverage_note") or "")}</p>')
    return head or f"{period}, {u['total_requests']:,} requests.", body


BUILDERS: Dict[str, Callable[[PackSource], Any]] = {
    "summary": _summary,
    "health_score": _health_score,
    "savings": _savings,
    "risk_distribution": _risk_distribution,
    "top_anomalies": _top_anomalies,
    "license_fit": _license_fit,
    "compliance": _compliance,
    "critical_objects": _critical_objects,
    "code_coverage": _code_coverage,
    "api_usage": _api_usage,
}


def build_visual(key: str, src: PackSource) -> Tuple[Optional[Body], Optional[str]]:
    """(caption, body) for a visual, or (None, note) when it's skipped."""
    v = VISUAL_BY_KEY[key]
    if v.user_level and not src.allow_users:
        return None, USERS_HIDDEN
    out = BUILDERS[key](src)
    return (None, out) if isinstance(out, str) else (out, None)


def availability(src: PackSource) -> List[Dict[str, Any]]:
    rows = []
    for v in VISUALS:
        body, note = build_visual(v.key, src)
        rows.append({"key": v.key, "title": v.title, "available": body is not None, "note": note,
                     "user_level": v.user_level})
    return rows


# ---------------------------------------------------------------- pages

STYLE = """
  body { margin: 0; font-family: 'Helvetica Neue', Arial, 'DejaVu Sans', sans-serif; color: #0b0b0b; font-size: 9pt; line-height: 1.4; }
  .frame { position: relative; box-sizing: border-box; padding: 30px 40px 26px; width: 800px; height: 450px; }
  .solid .frame, .solid .crop { background: #ffffff; }
  h1 { font-size: 20pt; margin: 0 0 4px; color: ACCENT; font-weight: 700; }
  .caption { color: #52514e; font-size: 8.5pt; margin: 0 0 14px; max-width: 600px; }
  .foot { position: absolute; left: 40px; right: 40px; bottom: 18px; display: flex; justify-content: space-between; font-size: 7pt; color: #898781; }
  .crop { box-sizing: border-box; width: 680px; padding: 12px; }
  .row { display: flex; gap: 28px; align-items: flex-start; }
  .grow { flex: 1; min-width: 0; }
  .hero { min-width: 230px; }
  .hero-value { font-size: 64pt; font-weight: 700; line-height: 1; }
  .hero-band { font-size: 14pt; font-weight: 600; margin: 4px 0 10px; }
  .of { font-size: 0.4em; color: #898781; font-weight: 400; }
  .big { font-size: 26pt; font-weight: 700; margin: 0 0 10px; }
  .unit { font-size: 10pt; color: #52514e; font-weight: 400; }
  .lead { font-size: 13pt; font-weight: 600; margin: 0 0 10px; }
  .label { font-size: 8pt; color: #52514e; font-weight: 600; margin: 0 0 4px; }
  .note { font-size: 7.5pt; color: #52514e; margin: 6px 0 0; }
  .tiles { display: flex; gap: 10px; margin: 12px 0; }
  .stat { flex: 1 1 0; border: 1px solid #e1e0d9; border-radius: 8px; padding: 6px 10px; background: #fcfcfb; }
  .stat-label { font-size: 7.5pt; color: #52514e; }
  .stat-value { font-size: 15pt; font-weight: 600; }
  table.data { width: 100%; border-collapse: collapse; font-size: 8.5pt; }
  table.data th { text-align: left; color: #52514e; font-weight: 600; border-bottom: 1px solid #c3c2b7; padding: 4px 6px; }
  table.data.compact td, table.data.compact th { padding: 2px 6px; }
  table.data td { border-bottom: 1px solid #e1e0d9; padding: 4px 6px; vertical-align: middle; }
  table.data .num { text-align: right; font-variant-numeric: tabular-nums; white-space: nowrap; }
  table.data .reason { color: #52514e; font-size: 7.5pt; }
  .grid { display: flex; flex-wrap: wrap; gap: 8px; }
  .tile { box-sizing: border-box; width: calc((100% - 16px) / 3); height: 98px; overflow: hidden; border: 1px solid #e1e0d9;
          border-radius: 8px; padding: 7px 10px; background: #fcfcfb; }
  .tile-label { font-size: 7.5pt; color: #52514e; font-weight: 600; }
  .tile-value { font-size: 19pt; font-weight: 700; line-height: 1.15; margin: 2px 0 1px; }
  .tile-note { font-size: 7.5pt; color: #52514e; margin: 0; }
  .tile-missing { font-size: 8pt; color: #898781; margin: 8px 0 0; }
  .strip { display: flex; gap: 2px; height: 7px; margin: 4px 0 4px; }
  .strip span { display: block; height: 7px; border-radius: 2px; }
  .keys { font-size: 7pt; color: #0b0b0b; }
  .key { margin-right: 7px; white-space: nowrap; }
  .key i { display: inline-block; width: 7px; height: 7px; border-radius: 2px; margin-right: 3px; }
  ol.top, ul.list { margin: 3px 0 0; padding: 0; list-style: none; font-size: 7.5pt; line-height: 1.3; }
  ol.top li, ul.list li { white-space: nowrap; overflow: hidden; text-overflow: ellipsis; }
  ol.top span { color: #52514e; }
"""


@dataclass
class Brand:
    """The one branding input images use: the title colour (the PDF's
    BrandContext also has it). Logos stay on the PDF letterhead only."""
    accent_hex: Optional[str] = None


def page_html(key: str, src: PackSource, brand: Any, opts: ImageOptions) -> str:
    body_parts, note = build_visual(key, src)
    if body_parts is None:
        raise ValueError(note or "not available")
    caption, body = body_parts
    accent = getattr(brand, "accent_hex", None) or "#14532d"
    title = VISUAL_BY_KEY[key].title
    solid = "" if opts.transparent else "solid"
    if opts.layout == "crop":
        w, h = CROP_W, 2000
        content = f'<div class="crop">{body}</div>'
    else:
        w, h = SLIDE_W, SLIDE_H
        content = (
            f'<div class="frame"><h1>{_esc(title)}</h1><p class="caption">{_esc(caption)}</p>{body}'
            + (f'<div class="foot"><span>{_esc(src.org_name)}</span></div>' if opts.org_name else "")
            + '</div>'
        )
    style = STYLE.replace("ACCENT", accent)
    return (f'<!doctype html><html><head><meta charset="utf-8"><style>@page {{ size: {w}px {h}px; margin: 0 }}{style}'
            f'</style></head><body class="{solid}">{content}</body></html>')


def _rasterise(html: str, opts: ImageOptions):
    """HTML -> RGBA PIL image at SCALE, cropped for the crop layout."""
    import pypdfium2 as pdfium
    from weasyprint import HTML  # type: ignore

    pdf = HTML(string=html).write_pdf()
    doc = pdfium.PdfDocument(pdf)
    try:
        page = doc[0]
        bitmap = page.render(scale=SCALE * PX_PER_PT, fill_color=(255, 255, 255, 0))  # px per pt
        img = bitmap.to_pil().convert("RGBA")
    finally:
        doc.close()
    if opts.layout == "crop":
        box = img.getchannel("A").getbbox()
        if box:
            pad = int(CROP_PAD * SCALE)
            left, top, right, bottom = box
            img = img.crop((max(0, left - pad), max(0, top - pad), min(img.width, right + pad), min(img.height, bottom + pad)))
    return img


def render_visual(key: str, src: PackSource, brand: Any, opts: ImageOptions) -> bytes:
    """One visual as PNG or JPEG bytes."""
    from PIL import Image

    img = _rasterise(page_html(key, src, brand, opts), opts)
    out = io.BytesIO()
    if opts.fmt == "png" and opts.transparent:
        img.save(out, "PNG", optimize=True)
    else:
        flat = Image.new("RGB", img.size, (255, 255, 255))
        flat.paste(img, mask=img.getchannel("A"))
        if opts.fmt == "jpeg":
            flat.save(out, "JPEG", quality=92, dpi=(192, 192))
        else:
            flat.save(out, "PNG", optimize=True)
    return out.getvalue()


def file_name(key: str, opts: ImageOptions, is_demo: bool) -> str:
    ext = "jpg" if opts.fmt == "jpeg" else "png"
    return f"{'demo-' if is_demo else ''}{key.replace('_', '-')}-{opts.layout}.{ext}"


def build_pptx(keys: List[str], src: PackSource, brand: Any, opts: ImageOptions) -> bytes:
    """One 16:9 slide per visual, the slide-layout image filling it."""
    from pptx import Presentation
    from pptx.util import Emu

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)  # 13.333 x 7.5 in
    blank = prs.slide_layouts[6]
    slide_opts = ImageOptions(fmt="png", transparent=False, layout="slide", org_name=opts.org_name)
    for key in keys:
        slide = prs.slides.add_slide(blank)
        img = render_visual(key, src, brand, slide_opts)
        slide.shapes.add_picture(io.BytesIO(img), 0, 0, width=prs.slide_width, height=prs.slide_height)
        slide.notes_slide.notes_text_frame.text = (
            f"{VISUAL_BY_KEY[key].title}, {src.org_name}." + (" Demo data: synthetic org." if src.is_demo else "")
        )
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def build_pack(keys: List[str], src: PackSource, brand: Any, opts: ImageOptions,
               layouts: List[str], images: bool = True, pptx: bool = False) -> Tuple[bytes, str, List[Dict[str, str]]]:
    """A ZIP of the chosen visuals in each layout (and the .pptx inside it
    when both are asked for), or the .pptx alone. Visuals that aren't
    available are skipped and listed in skipped.txt.
    Returns (bytes, media type, skipped)."""
    keys = [k for k in dict.fromkeys(keys) if k in VISUAL_BY_KEY]
    ready, skipped = [], []
    for k in keys:
        body, note = build_visual(k, src)
        (ready.append(k) if body else skipped.append({"key": k, "title": VISUAL_BY_KEY[k].title, "note": note or ""}))
    if not ready:
        raise ValueError("None of the chosen visuals is available.")
    deck = build_pptx(ready, src, brand, opts) if pptx else None
    if deck is not None and not images:
        return deck, "application/vnd.openxmlformats-officedocument.presentationml.presentation", skipped
    out = io.BytesIO()
    prefix = "demo-" if src.is_demo else ""
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for k in ready:
            for layout in layouts:
                o = ImageOptions(fmt=opts.fmt, transparent=opts.transparent, layout=layout, org_name=opts.org_name)
                z.writestr(file_name(k, o, src.is_demo), render_visual(k, src, brand, o))
        if deck is not None:
            z.writestr(f"{prefix}visuals.pptx", deck)
        if skipped:
            z.writestr("skipped.txt", "".join(f"{s['title']}: {s['note']}\n" for s in skipped))
        if src.is_demo:
            z.writestr("DEMO-DATA.txt", "These images show a synthetic demo org, not a real Salesforce org.\n")
    return out.getvalue(), "application/zip", skipped


async def gather_pack_source(db: Any, org_id: str) -> PackSource:
    """What the visuals draw from: the latest Health Report (active
    findings only, as in the PDF), the report context and the latest Data
    Quality run's critical objects. No Salesforce calls."""
    from types import SimpleNamespace

    from sqlalchemy import desc, select

    from app.domain.models import DataQualityRun, OrgAnalysisSnapshot, OrgFinding, Organization
    from app.services import privacy_mode
    from app.services.org_analyzer import OrgAnalyzerService
    from app.services.report.context import gather_report_context

    org = await db.get(Organization, org_id)
    snap = (await db.execute(
        select(OrgAnalysisSnapshot).where(OrgAnalysisSnapshot.organization_id == org_id)
        .order_by(desc(OrgAnalysisSnapshot.snapshot_at)).limit(1)
    )).scalar_one_or_none()
    findings = []
    if snap is not None:
        findings = list((await db.execute(select(OrgFinding).where(
            OrgFinding.organization_id == org_id, OrgFinding.snapshot_id == snap.id, OrgFinding.is_ignored.is_(False),
        ))).scalars().all())
    ctx = await gather_report_context(db, org_id, snap or SimpleNamespace(snapshot_at=datetime.now(timezone.utc)))
    dq_run = (await db.execute(
        select(DataQualityRun.id).where(DataQualityRun.organization_id == org_id).limit(1)
    )).scalar_one_or_none()
    return PackSource(
        org_name=(org.name if org and org.name else "Salesforce org"),
        is_demo=bool(org and org.is_demo),
        allow_users=privacy_mode.allows(org, privacy_mode.USER_RECORDS),
        snapshot=snap,
        findings=findings,
        ctx=ctx,
        critical_objects=await OrgAnalyzerService(db, org_id)._latest_critical_objects() if dq_run else [],
        data_quality_run=dq_run is not None,
    )
