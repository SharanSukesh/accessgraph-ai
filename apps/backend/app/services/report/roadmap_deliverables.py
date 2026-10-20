"""Roadmap deliverables (P1 part 2): closing slides and an SOW appendix.

Both are drawn from the same computed roadmap as the page and the Excel
(routes/roadmap._full), so the numbers match everywhere.

- **Closing slides**: a 16:9 PowerPoint rendered with the image pack's
  renderer and styles (item 3): the roadmap at a glance, the packages by
  phase, one slide per package (what was found, what the work is, effort
  and price) and the assumptions. Like the image pack, slides carry no
  logos; demo decks say so in each slide's notes and the file name.
- **SOW appendix**: an A4 PDF with the scope, approach, prerequisites,
  risks, effort and price of each package, the rate card and the
  assumptions, under the firm's name and accent colour.

Effort and prices are always ranges. Inferred and confirmed effort are
labelled as such; savings are labelled estimated; packages priced from
example playbooks say so.
"""
from __future__ import annotations

import io
from datetime import datetime, timezone
from typing import Any, Dict, List, Optional, Tuple

from app.services.report.image_pack import STYLE, ImageOptions, _esc, _rasterise

SEV_COLOURS = {"critical": "#7f1d1d", "high": "#b91c1c", "medium": "#b45309", "low": "#52514e", "info": "#898781"}
SLIDE_EXTRA = """
  .phases { display: flex; gap: 10px; margin-top: 10px; }
  .phase { flex: 1 1 0; border: 1px solid #e1e0d9; border-radius: 8px; padding: 7px 10px; background: #fcfcfb; min-width: 0; }
  .phase h2 { font-size: 9pt; margin: 0 0 4px; color: ACCENT; }
  .phase p { margin: 0 0 4px; font-size: 7.5pt; line-height: 1.3; }
  .phase b { font-weight: 600; }
  .muted { color: #52514e; }
  .cols { display: flex; gap: 24px; }
  .cols > div { flex: 1 1 0; min-width: 0; }
  ul.tight { margin: 2px 0 0; padding-left: 14px; font-size: 8pt; line-height: 1.35; }
  ul.found { list-style: none; padding-left: 0; }
  .dot { display: inline-block; width: 7px; height: 7px; border-radius: 2px; margin-right: 4px; }
"""


def _days(r: Optional[Dict[str, Any]]) -> str:
    return f"{r['low']:g}–{r['high']:g} days" if r else "not estimated"


def _money(n: Optional[float], cur: str) -> str:
    if n is None:
        return "—"
    sym = {"USD": "$", "GBP": "£", "EUR": "€"}.get(cur, "")
    return f"{sym}{n:,.0f}" if sym else f"{n:,.0f} {cur}"


def _price(r: Optional[Dict[str, Any]], cur: str) -> str:
    return f"{_money(r['low'], cur)}–{_money(r['high'], cur)}" if r else "not priced"


def _basis(p: Dict[str, Any]) -> str:
    if not p.get("effort"):
        return "not estimated"
    return f"{_days(p['effort'])}, {p['basis']}"


def _included(data: Dict[str, Any]) -> List[Dict[str, Any]]:
    return [p for p in data["packages"] if p.get("included", True)]


def assumptions(data: Dict[str, Any]) -> List[str]:
    t = data["totals"]
    out = [data["labels"]["effort"], data["labels"]["savings"]]
    if any(p["playbookExample"] for p in _included(data)):
        out.append("Some effort ranges come from example playbooks that haven't been reviewed for this engagement.")
    if t["unestimated"]:
        out.append(f"Not estimated yet: {', '.join(t['unestimated'])}.")
    if t["unpriced"]:
        out.append(f"Not priced (no day rate for a role): {', '.join(t['unpriced'])}.")
    if data["rateOverrides"]:
        out.append("Some day rates are specific to this engagement.")
    out.append("Prices exclude Salesforce licences and third-party costs.")
    if data.get("labels", {}).get("demo"):
        out.append(data["labels"]["demo"])
    return out


# ---------------------------------------------------------------- slides


def slides(data: Dict[str, Any]) -> List[Tuple[str, str, str]]:
    """(title, caption, body html) for each closing slide."""
    cur, t = data["currency"], data["totals"]
    inc = _included(data)
    split = []
    if t["inferred"]["packages"]:
        split.append(f"inferred {_days(t['inferred'])}")
    if t["confirmed"]["packages"]:
        split.append(f"confirmed {_days(t['confirmed'])}")
    tiles = "".join(
        f'<div class="stat"><div class="stat-label">{_esc(label)}</div><div class="stat-value">{_esc(value)}</div>'
        f'<div class="stat-label">{_esc(note)}</div></div>'
        for label, value, note in (
            ("Effort", _days(t["effort"]), "; ".join(split)),
            ("Price", _price(t["price"], cur) if t["price"]["high"] else "not priced", "from the rate card"),
            ("Estimated savings", _money(t["savingsCents"] / 100, "USD"), "a year, estimated"),
            ("Work packages", f"{t['packages']} in {t['phases']} phases", ""),
        ))
    phases = sorted({p["phase"] for p in inc})
    cols = "".join(
        f'<div class="phase"><h2>Phase {ph}</h2>'
        + "".join(f'<p><b>{_esc(p["title"])}</b><br><span class="muted">{_esc(_basis(p))} · {_esc(_price(p["price"], cur))}</span></p>'
                  for p in inc if p["phase"] == ph)
        + "</div>" for ph in phases[:4])
    out = [("Roadmap", "Effort and prices are ranges. Savings are estimated from the Health Report.",
            f'<div class="tiles">{tiles}</div><div class="phases">{cols}</div>')]

    rows = "".join(
        f'<tr><td class="num">{p["phase"]}</td><td>{_esc(p["title"])}</td><td>{_esc(_basis(p))}</td>'
        f'<td class="num">{_esc(_price(p["price"], cur))}</td>'
        f'<td class="num">{_esc(_money(p["savingsCents"] / 100, "USD")) if p["savingsCents"] else "—"}</td></tr>'
        for p in inc[:12])
    out.append(("Work packages", "Each package's effort says whether it was inferred from a playbook or confirmed by a consultant.",
                '<table class="data"><tr><th class="num">Phase</th><th>Package</th><th>Effort</th><th class="num">Price</th>'
                f'<th class="num">Est. savings / yr</th></tr>{rows}</table>'))

    for p in inc:
        found = "".join(
            f'<li><span class="dot" style="background:{SEV_COLOURS.get(f.get("severity") or "", "#898781")}"></span>{_esc(f["title"])}</li>'
            for f in [f for f in p["findings"] if f["present"]][:8]) or '<li class="muted">A package of the consultant\'s own.</li>'
        steps = "".join(f"<li>{_esc(s)}</li>" for s in (p.get("steps") or [])) or '<li class="muted">Scope to be agreed.</li>'
        after = ", ".join(next((q["title"] for q in data["packages"] if q["id"] == a), a) for a in p.get("after") or [])
        facts = "".join(
            f'<div class="stat"><div class="stat-label">{_esc(label)}</div><div class="stat-value">{_esc(value)}</div></div>'
            for label, value in (("Effort", _basis(p)), ("Price", _price(p["price"], cur)),
                                 ("Estimated savings", _money(p["savingsCents"] / 100, "USD") + " a year" if p["savingsCents"] else "—")))
        caption = f"Phase {p['phase']}" + (f", after {after}" if after else "") + (
            ". Effort from an example playbook." if p["playbookExample"] and p["basis"] == "inferred" else ".")
        out.append((p["title"], caption,
                    f'<div class="cols"><div><p class="label">What we found</p><ul class="tight found">{found}</ul></div>'
                    f'<div><p class="label">What the work is</p><ul class="tight">{steps}</ul></div></div>'
                    f'<div class="tiles">{facts}</div>'))

    items = "".join(f"<li>{_esc(a)}</li>" for a in assumptions(data))
    out.append(("Assumptions", "What the figures rest on.", f'<ul class="tight" style="font-size:9pt">{items}</ul>'))
    return out


def slide_html(title: str, caption: str, body: str, accent: str, org_name: Optional[str] = None) -> str:
    style = (STYLE + SLIDE_EXTRA).replace("ACCENT", accent)
    foot = f'<div class="foot"><span>{_esc(org_name)}</span></div>' if org_name else ""
    return ('<!doctype html><html><head><meta charset="utf-8"><style>@page { size: 800px 450px; margin: 0 }'
            f'{style}</style></head><body class="solid"><div class="frame"><h1>{_esc(title)}</h1>'
            f'<p class="caption">{_esc(caption)}</p>{body}{foot}</div></body></html>')


def build_slides_pptx(data: Dict[str, Any], accent: Optional[str], org_name: str, is_demo: bool) -> bytes:
    from PIL import Image
    from pptx import Presentation
    from pptx.util import Emu

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    blank = prs.slide_layouts[6]
    opts = ImageOptions(fmt="png", transparent=False, layout="slide")
    for title, caption, body in slides(data):
        img = _rasterise(slide_html(title, caption, body, accent or "#14532d", org_name), opts)
        flat = Image.new("RGB", img.size, (255, 255, 255))
        flat.paste(img, mask=img.getchannel("A"))
        buf = io.BytesIO()
        flat.save(buf, "PNG", optimize=True)
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(io.BytesIO(buf.getvalue()), 0, 0, width=prs.slide_width, height=prs.slide_height)
        slide.notes_slide.notes_text_frame.text = f"{title}, {org_name}." + (" Demo data: synthetic org." if is_demo else "")
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


# ---------------------------------------------------------------- SOW appendix

SOW_STYLE = """
  @page { size: A4; margin: 18mm 16mm 20mm; @bottom-right { content: "Page " counter(page) " of " counter(pages); font-family: 'Helvetica Neue', Arial, 'DejaVu Sans', sans-serif; font-size: 7.5pt; color: #898781; } }
  body { font-family: 'Helvetica Neue', Arial, 'DejaVu Sans', sans-serif; color: #0b0b0b; font-size: 9pt; line-height: 1.45; }
  h1 { font-size: 17pt; color: ACCENT; margin: 0 0 2px; }
  h2 { font-size: 12pt; color: ACCENT; margin: 18px 0 6px; border-bottom: 1px solid #e1e0d9; padding-bottom: 3px; }
  h3 { font-size: 10pt; margin: 12px 0 3px; }
  .meta { color: #52514e; font-size: 8.5pt; margin: 0 0 10px; }
  .demo { border: 1px solid #f59e0b; background: #fffbeb; color: #78350f; padding: 6px 8px; border-radius: 4px; font-size: 8.5pt; }
  table { width: 100%; border-collapse: collapse; font-size: 8.5pt; margin: 4px 0 8px; }
  th { text-align: left; color: #52514e; font-weight: 600; border-bottom: 1px solid #c3c2b7; padding: 3px 5px; }
  td { border-bottom: 1px solid #e1e0d9; padding: 3px 5px; vertical-align: top; }
  .num { text-align: right; white-space: nowrap; font-variant-numeric: tabular-nums; }
  .pkg { page-break-inside: avoid; }
  .label { font-weight: 600; color: #52514e; font-size: 8pt; margin: 6px 0 1px; }
  ul { margin: 0; padding-left: 15px; }
  .total td { font-weight: 600; }
"""


def sow_html(data: Dict[str, Any], org_name: str, firm_name: Optional[str], accent: Optional[str]) -> str:
    cur, t = data["currency"], data["totals"]
    inc = _included(data)
    today = datetime.now(timezone.utc).strftime("%d %B %Y")
    hr = (data.get("healthReport") or {}).get("at", "")[:10]

    def ul(items: List[str], empty: str) -> str:
        return "<ul>" + "".join(f"<li>{_esc(i)}</li>" for i in items) + "</ul>" if items else f'<p class="meta">{_esc(empty)}</p>'

    summary = "".join(
        f'<tr><td class="num">{p["phase"]}</td><td>{_esc(p["title"])}</td><td>{_esc(_basis(p))}</td>'
        f'<td class="num">{_esc(_price(p["price"], cur))}</td></tr>' for p in inc)
    summary += (f'<tr class="total"><td></td><td>Total</td><td>{_esc(_days(t["effort"]))}</td>'
                f'<td class="num">{_esc(_price(t["price"], cur)) if t["price"]["high"] else "not priced"}</td></tr>')
    pkgs = []
    for p in inc:
        effort = [f"Inferred: {_days(p['inferred'])}" + (f" (sized by {p['inferred']['size']} {p['inferred'].get('itemLabel') or 'items'})"
                                                         if p.get("inferred") else "")]
        if p.get("confirmed"):
            effort.append(f"Confirmed: {_days(p['confirmed'])}")
        effort.append(f"Used: {_basis(p)}")
        pb = f"{p['playbookTitle']} {p['playbookVersion']}" + (" (example values)" if p["playbookExample"] else "") if p["playbookTitle"] else "none"
        pkgs.append(
            f'<div class="pkg"><h3>{_esc(p["title"])}</h3>'
            f'<p class="meta">Phase {p["phase"]} · playbook {_esc(pb)}</p>'
            f'<p class="label">Scope: findings addressed</p>{ul([f["title"] for f in p["findings"] if f["present"]], "Defined by the consultant.")}'
            f'<p class="label">Approach</p>{ul(p.get("steps") or [], "To be agreed.")}'
            f'<p class="label">Prerequisites</p>{ul(p.get("prerequisites") or [], "None listed.")}'
            f'<p class="label">Risks</p>{ul(p.get("risks") or [], "None listed.")}'
            f'<p class="label">Effort and price</p>{ul(effort + [f"Price: {_price(p["price"], cur)}"], "")}'
            + (f'<p class="label">Notes</p><p>{_esc(p["notes"])}</p>' if p.get("notes") else "")
            + "</div>")
    rates = "".join(f'<tr><td>{_esc(r["role"])}</td><td class="num">{_esc(_money(r["dayRate"], cur))}</td>'
                    f'<td>{"This engagement" if r["override"] else "Rate card"}</td></tr>' for r in data["rates"])
    left_out = [p["title"] for p in data["packages"] if not p.get("included", True)]
    body = (
        f'<h1>Appendix: scope, effort and pricing</h1>'
        f'<p class="meta">{_esc(data["name"])}' + ("" if data["name"].startswith(org_name) else f" · {_esc(org_name)}") + (f" · prepared by {_esc(firm_name)}" if firm_name else "")
        + f" · {today}" + (f" · Health Report of {hr}" if hr else "") + "</p>"
        + (f'<p class="demo">{_esc(data["labels"]["demo"])}</p>' if data["labels"].get("demo") else "")
        + '<h2>Summary</h2><table><tr><th class="num">Phase</th><th>Work package</th><th>Effort</th><th class="num">Price</th></tr>'
        + summary + "</table>"
        + f'<p>Estimated annual savings from the findings addressed: {_esc(_money(t["savingsCents"] / 100, "USD"))} (estimated).</p>'
        + (f'<p class="meta">Out of scope for now: {_esc(", ".join(left_out))}.</p>' if left_out else "")
        + "<h2>Work packages</h2>" + "".join(pkgs)
        + '<h2>Day rates</h2><table><tr><th>Role</th><th class="num">Day rate</th><th>Source</th></tr>'
        + (rates or '<tr><td colspan="3">No rate card set.</td></tr>') + "</table>"
        + "<h2>Assumptions</h2>" + ul(assumptions(data), "")
    )
    style = SOW_STYLE.replace("ACCENT", accent or "#14532d")
    return (f'<!doctype html><html><head><meta charset="utf-8"><title>{_esc(data["name"])}: SOW appendix</title>'
            f"<style>{style}</style></head><body>{body}</body></html>")


def build_sow_pdf(data: Dict[str, Any], org_name: str, firm_name: Optional[str], accent: Optional[str]) -> bytes:
    from weasyprint import HTML  # type: ignore

    return HTML(string=sow_html(data, org_name, firm_name, accent)).write_pdf()
