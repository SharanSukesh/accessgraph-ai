"""Excel workbook of a priced roadmap (P1).

Sheets: Summary, Roadmap (one row per package: phase, effort with its
basis, price, estimated savings, dependencies), Findings (each package's
findings), Rate card, Playbooks. Effort and prices are ranges; inferred
and consultant-confirmed effort are separate columns, never merged.
"""
from __future__ import annotations

import io
from typing import Any, Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from app.domain.models import Organization
from app.services.report.workbook import USD, Column, _safe, _table

DAYS = '0.0" d"'


def build_roadmap_workbook(org: Organization, data: Dict[str, Any], playbooks: List[Dict[str, Any]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    t = data["totals"]
    packages = data["packages"]
    titles = {p["id"]: p["title"] for p in packages}
    rows: List[tuple] = [
        ("Roadmap", data["name"], None),
        ("Organization", org.name, None),
        ("Health Report", (data.get("healthReport") or {}).get("at", "")[:10] or None, None),
        ("Packages in plan", t["packages"], "0"),
        ("Phases", t["phases"], "0"),
        ("Effort, low (days)", t["effort"]["low"], DAYS),
        ("Effort, high (days)", t["effort"]["high"], DAYS),
        ("  of which inferred", f"{t['inferred']['low']:g}-{t['inferred']['high']:g} days, {t['inferred']['packages']} packages", None),
        ("  of which confirmed", f"{t['confirmed']['low']:g}-{t['confirmed']['high']:g} days, {t['confirmed']['packages']} packages", None),
        ("Price, low", t["price"]["low"], USD),
        ("Price, high", t["price"]["high"], USD),
        ("Estimated annual savings", t["savingsCents"] / 100, USD),
        ("Not estimated", ", ".join(t["unestimated"]) or "None", None),
        ("Not priced", ", ".join(t["unpriced"]) or "None", None),
        ("Effort", data["labels"]["effort"], None),
        ("Savings", data["labels"]["savings"], None),
    ]
    if data.get("isDemo"):
        rows.insert(0, ("Demo data", data["labels"]["demo"], None))
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 90
    for r, (label, value, fmt) in enumerate(rows, start=1):
        ws.cell(row=r, column=1, value=label).font = Font(bold=True)
        cell = ws.cell(row=r, column=2, value=_safe(value))
        cell.alignment = Alignment(horizontal="left", wrap_text=True, vertical="top")
        if fmt:
            cell.number_format = fmt

    def rng(x: Any, end: str) -> Any:
        return x[end] if x else None

    _table(wb.create_sheet("Roadmap"), [
        Column("Phase", 8, "0"), Column("Package", 36, wrap=True), Column("In plan", 9), Column("Findings", 9, "0"),
        Column("Inferred low", 12, DAYS), Column("Inferred high", 12, DAYS), Column("Confirmed low", 13, DAYS),
        Column("Confirmed high", 13, DAYS), Column("Effort used", 12), Column("Day rate (blended)", 16, USD),
        Column("Price low", 13, USD), Column("Price high", 13, USD), Column("Estimated savings / yr", 18, USD),
        Column("After", 30, wrap=True), Column("Playbook", 30, wrap=True), Column("Notes", 50, wrap=True),
    ], ([p["phase"], p["title"], "Yes" if p.get("included", True) else "No", sum(f["present"] for f in p["findings"]),
         rng(p["inferred"], "low"), rng(p["inferred"], "high"), rng(p.get("confirmed"), "low"), rng(p.get("confirmed"), "high"),
         p["basis"] or "not estimated", p["blendedRate"], rng(p["price"], "low"), rng(p["price"], "high"),
         p["savingsCents"] / 100 if p["savingsCents"] else None, ", ".join(titles.get(a, a) for a in p.get("after") or []),
         (f"{p['playbookTitle']} ({p['playbookVersion']}{', example' if p['playbookExample'] else ''})"
          if p["playbookTitle"] else None), p.get("notes") or None] for p in packages), "No packages.")

    _table(wb.create_sheet("Findings"), [
        Column("Package", 32, wrap=True), Column("Finding", 60, wrap=True), Column("Severity", 10), Column("Affected", 10, "0"),
        Column("Estimated savings / yr", 18, USD), Column("Review", 12),
    ], ([p["title"], f["title"], f.get("severity"), f.get("affected"),
         (f.get("savingsCents") or 0) / 100 if f.get("savingsCents") else None, f.get("review")]
        for p in packages for f in p["findings"]), "No findings.")

    _table(wb.create_sheet("Rate card"), [Column("Role", 30), Column("Day rate", 14, USD), Column("Source", 26)],
           ([r["role"], r["dayRate"], "Engagement override" if r["override"] else
             ("Demo rate card" if data.get("isDemo") else "Firm rate card")] for r in data["rates"]), "No rate card set.")

    _table(wb.create_sheet("Playbooks"), [
        Column("Playbook", 32, wrap=True), Column("Version", 12), Column("Example values", 14), Column("Finding codes", 36, wrap=True),
        Column("Typical days", 14), Column("Per item", 16), Column("Role mix", 36, wrap=True), Column("Steps", 60, wrap=True),
        Column("Prerequisites", 40, wrap=True), Column("Risks", 40, wrap=True),
    ], ([pb["title"], pb.get("version"), "Yes" if pb.get("example") else "No", ", ".join(pb.get("codes") or []),
         f"{pb['baseDays'][0]:g}-{pb['baseDays'][1]:g}",
         f"{pb['perItemDays'][0]:g}-{pb['perItemDays'][1]:g} each ({pb.get('itemLabel') or 'items'})" if any(pb["perItemDays"]) else None,
         ", ".join(f"{r} {round(s * 100)}%" for r, s in (pb.get("roles") or {}).items()), "\n".join(pb.get("steps") or []),
         "\n".join(pb.get("prerequisites") or []), "\n".join(pb.get("risks") or [])] for pb in playbooks), "No playbooks.")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
