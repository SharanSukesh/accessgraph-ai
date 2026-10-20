"""Org documentation, item 11 part c: flow detail (#9), the field x layout
matrix (#4), approval steps (#18), external data sources and record types
through permission sets (#11)."""
import io

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.domain.models import (
    PermissionSetAssignmentSnapshot,
    PermissionSetGroupComponentSnapshot,
    ProfileSnapshot,
    UserSnapshot,
)
from app.services import metadata_detail as md
from app.services import ui_access as ua
from app.services.effective_access import EffectiveAccessService
from tests.test_demo_orgs import client  # noqa: F401
from tests.test_demo_path import lumenvale  # noqa: F401
from tests.test_documentation_export import _rows

FINANCE_REPORTING = "0PS000000000022DMO"


def test_flow_detail_reads_elements_and_record_operations():
    d = md.flow_detail({
        "start": {"object": "Opportunity", "recordTriggerType": "Update", "filters": [{"field": "StageName"}]},
        "variables": [{"name": "v1"}, {"name": "v2"}],
        "decisions": {"name": "only_one"},  # a single element comes back as an object, not a list
        "recordLookups": [{"name": "Get_Account", "label": "Get Account", "object": "Account", "queriedFields": ["Rating", "Name"],
                           "filters": [{"field": "Id"}]}],
        "recordUpdates": [{"name": "Update_Record", "inputReference": "$Record", "inputAssignments": [{"field": "NextStep"}]},
                          {"name": "Save_All", "inputReference": "Changed"}],
        "subflows": [{"flowName": "Update_Territory"}],
        "actionCalls": [{"actionName": "submit", "actionType": "submit"}],
    })
    assert d["elements"] == {"variables": 2, "decisions": 1, "recordLookups": 1, "recordUpdates": 2, "subflows": 1,
                             "actionCalls": 1}
    assert (d["start_object"], d["record_trigger"]) == ("Opportunity", "Update")
    ops = {o["element"]: o for o in d["operations"]}
    assert ops["Start"]["filter_fields"] == ["StageName"] and ops["Start"]["operation"] == "Trigger condition"
    assert ops["Get Account"]["fields"] == ["Name", "Rating"] and ops["Get Account"]["filter_fields"] == ["Id"]
    # Updating the record that started the flow: the start object.
    assert (ops["Update_Record"]["object"], ops["Update_Record"]["fields"]) == ("Opportunity", ["NextStep"])
    assert ops["Save_All"]["object"] is None and ops["Save_All"]["from_variable"] == "Changed"
    assert d["subflows"] == ["Update_Territory"] and d["actions"] == [{"name": "submit", "type": "submit"}]
    assert md.flow_detail(None) is None


def test_layout_fields_and_set_record_types():
    placed = md.layout_fields({"layoutSections": [
        {"label": "Info", "layoutColumns": [{"layoutItems": [{"field": "Name", "behavior": "Required"},
                                                             {"field": "Tax_ID__c", "behavior": "Readonly"}]},
                                            {"layoutItems": {"field": "Name", "behavior": "Edit"}}]},  # listed once
        {"label": "Custom Links", "layoutColumns": [{"layoutItems": [{"customLink": "Open_in_ERP"}]}]},
    ]})
    assert placed == [{"field": "Name", "behavior": "Required", "section": "Info"},
                      {"field": "Tax_ID__c", "behavior": "Readonly", "section": "Info"}]
    assert md.set_record_types({"recordTypeVisibilities": [
        {"recordType": "Opportunity.Partner_Deal", "visible": True},
        {"recordType": "Opportunity.Renewal", "visible": False}]}) == {"Opportunity": ["Opportunity.Partner_Deal"]}


@pytest.mark.asyncio
async def test_live_reads_layout_and_set_metadata_and_stop_on_an_unverified_path():
    queries = []

    class Fake:
        async def query_all(self, q):
            queries.append(q)
            if "FROM PermissionSet WHERE IsOwnedByProfile = false" in q:
                return [{"Id": "0PSA"}, {"Id": "0PSB"}, {"Id": "0PSC"}]
            if "FlowDefinitionView" in q:
                return [{"DurableId": "300A", "ApiName": "Sync", "Label": "Sync", "ProcessType": "AutoLaunchedFlow",
                         "TriggerType": "RecordAfterSave", "IsActive": True, "ActiveVersionId": "301A"}]
            raise RuntimeError("not available")

        async def query_tooling(self, q):
            queries.append(q)
            if "FROM Layout WHERE" in q:
                return [{"Metadata": {"layoutSections": [{"label": "S", "layoutColumns": [
                    {"layoutItems": [{"field": "Name", "behavior": "Required"}]}]}]}}]
            if "FROM Layout" in q:
                return [{"Id": "00hA", "Name": "Account Layout", "TableEnumOrId": "Account"}]
            if "FROM Flow WHERE" in q:
                return [{"Metadata": {"start": {"object": "Account"}, "recordUpdates": [
                    {"name": "u", "object": "Account", "inputAssignments": [{"field": "Rating"}]}]}}]
            if "FROM PermissionSet WHERE Id" in q:
                raise RuntimeError("sObject type 'PermissionSet' is not supported")
            raise RuntimeError("not available")

        async def describe_object(self, name):
            raise RuntimeError("no describe")

    out = await ua.collect_live(Fake())
    assert out["layout_fields"] == {"00hA": [{"field": "Name", "behavior": "Required", "section": "S"}]}
    assert out["flow_details"]["300A"]["operations"][0]["fields"] == ["Rating"]
    # The set metadata read failed once and wasn't repeated for the other sets.
    assert sum("FROM PermissionSet WHERE Id" in q for q in queries) == 1 and "set_record_types" not in out
    assert any("Record types granted by permission sets couldn't be read" in n for n in out["notes"])
    assert "'ExternalDataSource'" in next(q for q in queries if "SetupEntityAccess" in q)
    assert all(q.lstrip().upper().startswith("SELECT") for q in queries)  # reads only


async def _active_users_by_profile(db, org_id, names):
    pids = {p.salesforce_id for p in (await db.execute(select(ProfileSnapshot).where(
        ProfileSnapshot.organization_id == org_id, ProfileSnapshot.name.in_(names)))).scalars().all()}
    return {u.salesforce_id for u in (await db.execute(select(UserSnapshot).where(
        UserSnapshot.organization_id == org_id, UserSnapshot.is_active.is_(True),
        UserSnapshot.profile_id.in_(pids)))).scalars().all()}


@pytest.mark.asyncio
async def test_part_c_sheets_on_lumenvale(client, async_db_session, lumenvale):  # noqa: F811
    db = async_db_session
    org_id, cookie = lumenvale
    wb = load_workbook(io.BytesIO((await client.get(f"/orgs/{org_id}/documentation.xlsx", cookies=cookie)).content))
    access = await EffectiveAccessService(db).org_access(org_id)

    def can_read(uid, key, kind="fields"):
        return bool(((access.get(uid) or {}).get(kind, {}).get(key) or {}).get("read"))

    # Fields on layouts: who gets the layout (by profile), and of them who can read the field.
    _, placed = _rows(wb["Fields on Layouts"])
    sales = await _active_users_by_profile(db, org_id, ["Sales Director", "Sales Manager", "Sales Rep"])
    tax = [r for r in placed if r["Field"] == "Tax_ID__c"]
    assert [r["Layout"] for r in tax] == ["Account - Sales Layout"] and tax[0]["Behavior"] == "Readonly"
    assert tax[0]["Layout assigned to profiles"] == "Sales Director, Sales Manager, Sales Rep"
    assert tax[0]["Active users who get the layout"] == len(sales)
    assert tax[0]["Of them, can read the field"] == sum(can_read(u, "Account.Tax_ID__c") for u in sales)
    every = {u for u in access}
    status = next(r for r in placed if r["Layout"] == "Case Layout" and r["Field"] == "Status")
    assert status["Active users who get the layout"] == len(every)
    assert status["Of them, can read the field"] == sum(can_read(u, "Case.Status") for u in every) < len(every)
    # Name can't have field-level security: it follows the object's access.
    name = next(r for r in placed if r["Layout"] == "Case Layout" and r["Field"] == "CaseNumber")
    assert name["Of them, can read the field"] == sum(can_read(u, "Case", "objects") for u in every)
    legacy = [r for r in placed if r["Layout"] == "Account - Legacy Layout"]
    assert legacy and all(r["Layout assigned to profiles"] == "Not assigned" for r in legacy)

    # Flow elements and the fields each flow reads or writes.
    head, elements = _rows(wb["Flow Elements"])
    assert len(elements) == 15 and len(head) == len(set(head))
    discount = next(r for r in elements if r["API name"] == "Discount_Approval_Request")
    assert (discount["Screens"], discount["Get records"], discount["Update records"]) == (2, 1, 1)
    assert discount["Actions called"] == "submit (submit)"
    sync = next(r for r in elements if r["API name"] == "Opportunity_Stage_Sync")
    assert (sync["Start object"], sync["Runs when"]) == ("Opportunity", "Updated")
    _, ffields = _rows(wb["Flow Fields"])
    assert {"Flow": "Opportunity Stage Sync", "Element": "Start", "Operation": "Trigger condition", "Object": "Opportunity",
            "Field": "StageName", "Used as": "Condition", "Note": None} in ffields
    written = {(r["Object"], r["Field"]) for r in ffields if r["Used as"] == "Written"}
    assert ("Opportunity", "Discount_Percent__c") in written and ("Account", "Rating") in written
    reminder = next(r for r in ffields if r["Element"] == "Create Reminder Tasks")
    assert reminder["Note"] == "Whole record from the variable ReminderTasks" and reminder["Field"] is None
    # Item 13's unreferenced cleanup candidate stays unreferenced.
    assert not any(r["Field"] == "Legacy_Region__c" for r in ffields + placed)

    # Approval steps: by name; entry criteria and approvers aren't read.
    _, steps = _rows(wb["Approval Steps"])
    assert [(s["Approval process"], s["Step"]) for s in steps][:2] == [
        ("Discount over 20%", "Sales Director Review"), ("Discount over 20%", "Finance Review over 40%")]
    assert len(steps) == 4 and all(s["Entry criteria and approvers"].startswith("Not read") for s in steps)

    # External data sources, with who is granted each through their profile, sets and groups.
    _, sources = _rows(wb["External Data Sources"])
    erp = next(s for s in sources if s["External data source"] == "ERP_Orders")
    assert erp["Access granted by"] == ("Finance Director (Profile), Finance Reporting (in Finance Ops Bundle), "
                                        "System Administrator (Profile)")
    granted = await _active_users_by_profile(db, org_id, ["Finance Director", "System Administrator"])
    groups = {c.permission_set_group_id for c in (await db.execute(select(PermissionSetGroupComponentSnapshot).where(
        PermissionSetGroupComponentSnapshot.organization_id == org_id,
        PermissionSetGroupComponentSnapshot.permission_set_id == FINANCE_REPORTING))).scalars().all()}
    holders = {a.assignee_id for a in (await db.execute(select(PermissionSetAssignmentSnapshot).where(
        PermissionSetAssignmentSnapshot.organization_id == org_id,
        PermissionSetAssignmentSnapshot.permission_set_id.in_({FINANCE_REPORTING} | groups)))).scalars().all()}
    assert erp["Active users granted it"] == len(granted | (holders & every))

    # Record types made available through permission sets.
    _, rts = _rows(wb["Record Types"])
    partner = next(r for r in rts if r["Record type"] == "Partner Deal")
    assert partner["Available through permission sets"] == "Opportunity Approvals (in Sales Manager Bundle)"
    about = {row[0]: row[1] for row in wb["About"].iter_rows(values_only=True)}
    assert "approvers" in about["Not included yet"] and "Flow detail" not in about["Not included yet"]
