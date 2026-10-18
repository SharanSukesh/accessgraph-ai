"""Image pack (roadmap item 3): every format and layout, sizes and
transparency, the demo label on every image, masked names never shown,
user-level visuals left out at Metadata only, and the ZIP / PPTX."""
import io
import zipfile

import pytest
import pytest_asyncio
from PIL import Image

from app.domain.models import Organization
from app.services import privacy_mode
from app.services.data_quality import DataQualityService
from app.services.org_analyzer import OrgAnalyzerService
from app.services.report import image_pack as ip
from tests.test_demo_orgs import client  # noqa: F401
from tests.test_demo_path import lumenvale  # noqa: F401


def _can_render() -> bool:
    try:
        from weasyprint import HTML  # type: ignore
        HTML(string="<p>x</p>").write_pdf()
        return True
    except (ImportError, OSError):
        return False


needs_render = pytest.mark.skipif(not _can_render(), reason="WeasyPrint system libraries not installed")


def _pdf_text(html: str) -> str:
    import pypdfium2 as pdfium
    from weasyprint import HTML  # type: ignore
    doc = pdfium.PdfDocument(HTML(string=html).write_pdf())
    try:
        page = doc[0]
        return page.get_textpage().get_text_range()
    finally:
        doc.close()


async def _analysed(db, org_id):
    await DataQualityService(db, org_id).run(actor_email="test")
    await OrgAnalyzerService(db, org_id).run(actor_email="test")
    return await ip.gather_pack_source(db, org_id)


@pytest.mark.asyncio
async def test_lumenvale_availability(async_db_session, lumenvale):  # noqa: F811
    src = await _analysed(async_db_session, lumenvale[0])
    avail = {v["key"]: v for v in ip.availability(src)}
    for key in ("health_score", "savings", "risk_distribution", "top_anomalies", "critical_objects",
                "code_coverage", "api_usage"):
        assert avail[key]["available"], (key, avail[key]["note"])
    # Modules that haven't run are skipped with a note, never drawn empty.
    assert not avail["license_fit"]["available"] and avail["license_fit"]["note"] == "Run License Fit first."
    assert not avail["compliance"]["available"] and "compliance" in avail["compliance"]["note"]


@needs_render
@pytest.mark.asyncio
async def test_every_format_and_layout(async_db_session, lumenvale):  # noqa: F811
    src = await _analysed(async_db_session, lumenvale[0])
    shots = {}
    for fmt, transparent, layout in [("png", False, "slide"), ("png", True, "slide"), ("jpeg", False, "slide"),
                                     ("png", False, "crop"), ("png", True, "crop"), ("jpeg", True, "crop")]:
        data = ip.render_visual("code_coverage", src, None, ip.ImageOptions(fmt=fmt, transparent=transparent, layout=layout))
        shots[(fmt, transparent, layout)] = Image.open(io.BytesIO(data))
    slide = shots[("png", False, "slide")]
    assert slide.size == (3840, 2160) and slide.mode == "RGB" and slide.getpixel((5, 5)) == (255, 255, 255)
    clear = shots[("png", True, "slide")]
    assert clear.size == (3840, 2160) and clear.mode == "RGBA" and clear.getpixel((5, 5))[3] == 0
    jpeg = shots[("jpeg", False, "slide")]
    assert jpeg.format == "JPEG" and jpeg.mode == "RGB" and jpeg.size == (3840, 2160)
    crop = shots[("png", False, "crop")]
    assert crop.width < 3840 and crop.height < 2160 and crop.width > 2000
    assert shots[("png", True, "crop")].getpixel((2, 2))[3] == 0
    # JPEG can't be transparent, whatever is asked.
    assert shots[("jpeg", True, "crop")].mode == "RGB"


@needs_render
@pytest.mark.asyncio
async def test_demo_label_on_every_image_whatever_the_toggles(async_db_session, lumenvale):  # noqa: F811
    src = await _analysed(async_db_session, lumenvale[0])
    for key in [v["key"] for v in ip.availability(src) if v["available"]]:
        for layout in ip.LAYOUTS:
            opts = ip.ImageOptions(layout=layout, firm_logo=False, client_logo=False, transparent=True)
            assert ip.DEMO_LABEL in _pdf_text(ip.page_html(key, src, None, opts)), (key, layout)
    src.is_demo = False
    assert ip.DEMO_LABEL not in _pdf_text(ip.page_html("health_score", src, None, ip.ImageOptions()))


@pytest_asyncio.fixture
async def masked_lumenvale(async_db_session, lumenvale):  # noqa: F811
    """Lumenvale switched to Masked: the purge aliases every stored user, as
    it does for a client org. (A demo org's sync loads the fixture as is.)"""
    from app.services.privacy_purge import enforce_privacy_mode
    org = await async_db_session.get(Organization, lumenvale[0])
    privacy_mode.set_privacy(org, privacy_mode.MASKED, True, "test")
    await enforce_privacy_mode(async_db_session, org)
    await async_db_session.flush()
    return lumenvale


@needs_render
@pytest.mark.asyncio
async def test_masked_names_never_appear(async_db_session, masked_lumenvale):
    src = await _analysed(async_db_session, masked_lumenvale[0])
    names = [a["name"] for a in src.ctx.flagged_access]
    assert names and all(n.startswith("User ") for n in names)
    text = _pdf_text(ip.page_html("top_anomalies", src, None, ip.ImageOptions()))
    assert names[0] in text
    # Real names from the fixture (the top-ranked anomalies on the unmasked org).
    for real in ("Mina Tanaka", "Hayden Galloway", "Ayla Zielinski"):
        assert real not in text
    api = _pdf_text(ip.page_html("api_usage", src, None, ip.ImageOptions()))
    assert "Ayla Zielinski" not in api


@pytest.mark.asyncio
async def test_metadata_only_leaves_user_level_visuals_out(async_db_session, lumenvale):  # noqa: F811
    db = async_db_session
    org = await db.get(Organization, lumenvale[0])
    await OrgAnalyzerService(db, lumenvale[0]).run(actor_email="test")
    privacy_mode.set_privacy(org, privacy_mode.METADATA_ONLY, False, "test")
    await db.flush()
    src = await ip.gather_pack_source(db, lumenvale[0])
    avail = {v["key"]: v for v in ip.availability(src)}
    for key in ("risk_distribution", "top_anomalies", "license_fit"):
        assert not avail[key]["available"] and avail[key]["note"] == ip.USERS_HIDDEN
    assert avail["health_score"]["available"]


@needs_render
@pytest.mark.asyncio
async def test_zip_and_pptx(async_db_session, lumenvale):  # noqa: F811
    from pptx import Presentation

    src = await _analysed(async_db_session, lumenvale[0])
    data, media, skipped = ip.build_pack(["health_score", "api_usage", "compliance"], src, None,
                                         ip.ImageOptions(fmt="png", transparent=True), ["slide", "crop"], pptx=True)
    assert media == "application/zip" and [s["key"] for s in skipped] == ["compliance"]
    z = zipfile.ZipFile(io.BytesIO(data))
    assert sorted(z.namelist()) == sorted([
        "demo-health-score-slide.png", "demo-health-score-crop.png", "demo-api-usage-slide.png",
        "demo-api-usage-crop.png", "demo-visuals.pptx", "skipped.txt", "DEMO-DATA.txt",
    ])
    assert "Compliance control indicators" in z.read("skipped.txt").decode()
    deck = Presentation(io.BytesIO(z.read("demo-visuals.pptx")))
    assert len(deck.slides) == 2 and (deck.slide_width, deck.slide_height) == (12192000, 6858000)
    assert "Demo data" in deck.slides[0].notes_slide.notes_text_frame.text

    alone, media, _ = ip.build_pack(["savings"], src, None, ip.ImageOptions(), [], images=False, pptx=True)
    assert media.endswith("presentationml.presentation") and len(Presentation(io.BytesIO(alone)).slides) == 1
    with pytest.raises(ValueError):
        ip.build_pack(["compliance"], src, None, ip.ImageOptions(), ["slide"])


@needs_render
@pytest.mark.asyncio
async def test_routes(client, async_db_session, lumenvale):  # noqa: F811
    org_id, cookie = lumenvale
    await _analysed(async_db_session, org_id)
    listing = (await client.get(f"/orgs/{org_id}/image-pack", cookies=cookie)).json()
    assert listing["is_demo"] and len(listing["visuals"]) == len(ip.VISUALS)

    r = await client.get(f"/orgs/{org_id}/image-pack/health_score", params={"format": "jpeg", "layout": "crop"}, cookies=cookie)
    assert r.status_code == 200 and r.headers["content-type"] == "image/jpeg"
    assert 'filename="demo-health-score-crop.jpg"' in r.headers["content-disposition"]
    r = await client.get(f"/orgs/{org_id}/image-pack/compliance", cookies=cookie)
    assert r.status_code == 409 and "compliance" in r.json()["detail"]
    assert (await client.get(f"/orgs/{org_id}/image-pack/nope", cookies=cookie)).status_code == 404
    assert (await client.get(f"/orgs/{org_id}/image-pack/health_score", params={"format": "gif"},
                             cookies=cookie)).status_code == 400

    r = await client.post(f"/orgs/{org_id}/image-pack", cookies=cookie,
                          json={"visuals": ["health_score"], "format": "png", "layouts": ["slide"]})
    assert r.status_code == 200 and r.headers["content-type"] == "application/zip"
    assert zipfile.ZipFile(io.BytesIO(r.content)).namelist()[0] == "demo-health-score-slide.png"
    # No access without a session.
    assert (await client.get(f"/orgs/{org_id}/image-pack")).status_code in (401, 403)
