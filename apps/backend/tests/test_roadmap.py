"""Priced roadmap builder (P1): packages from findings, inferred vs
confirmed effort, prices from the rate card, phases, the firm's data,
reviews, privacy and the workbook."""
import io

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.domain.models import FirmSetting, ItemReview, Organization, OrgAccessGrant, OrgUser, OrgUserRole
from app.services import privacy_mode
from app.services import roadmap as rm
from app.services.data_quality import DataQualityService
from app.services.org_analyzer import OrgAnalyzerService
from app.services.privacy_purge import enforce_privacy_mode
from tests.test_demo_orgs import _cookie, client  # noqa: F401
from tests.test_demo_path import lumenvale  # noqa: F401


async def _build(client, db, org_id, cookie):  # noqa: F811
    await DataQualityService(db, org_id).run(actor_email="test")  # its findings join the Health Report's
    await OrgAnalyzerService(db, org_id).run(actor_email="test")
    r = await client.post(f"/orgs/{org_id}/roadmaps", json={}, cookies=cookie)
    assert r.status_code == 200, r.text
    return r.json()


@pytest.mark.asyncio
async def test_lumenvale_roadmap(client, async_db_session, lumenvale):  # noqa: F811
    db = async_db_session
    org_id, cookie = lumenvale
    r = await _build(client, db, org_id, cookie)
    assert r["isDemo"] is True and r["labels"]["demo"] and r["name"] == "Lumenvale Systems (Demo) roadmap"
    by = {p["playbook"]: p for p in r["packages"]}
    assert set(by) == {"licence-reclaim", "licence-rightsizing", "permission-cleanup", "apex-coverage", "data-storage"}
    assert r["notInRoadmap"] == [] and sum(len(p["findings"]) for p in r["packages"]) == 18
    # Savings are the Health Report's, labelled estimated.
    assert r["totals"]["savingsCents"] == 18_318_000 and "estimated" in r["labels"]["savings"]
    # Inferred effort: base + per item x size, in half days. 27 users (11 + 16) for the reclaim.
    reclaim = by["licence-reclaim"]
    assert reclaim["inferred"] == {"low": 2.0, "high": 5.0, "size": 27, "itemLabel": "users", "capped": False}
    assert reclaim["basis"] == "inferred" and reclaim["blendedRate"] == 1200
    assert reclaim["price"] == {"low": 2400, "high": 6000}
    apex = by["apex-coverage"]  # 36 classes: 1 + 9 = 10, 2 + 27 = 29; developer 80% at 1100 + architect 20% at 1600
    assert (apex["inferred"]["low"], apex["inferred"]["high"], apex["blendedRate"]) == (10.0, 29.0, 1200)
    # Phases follow dependencies: right-sizing and permissions come after the reclaim.
    assert reclaim["phase"] == 1 and by["licence-rightsizing"]["phase"] == 2 and by["permission-cleanup"]["phase"] == 2
    assert by["licence-rightsizing"]["after"] == [reclaim["id"]]
    t = r["totals"]
    assert t["confirmed"]["packages"] == 0 and t["inferred"]["packages"] == 5 and t["phases"] == 2
    assert t["price"]["low"] == sum(p["price"]["low"] for p in r["packages"])


@pytest.mark.asyncio
async def test_confirmed_effort_stays_apart(client, async_db_session, lumenvale):  # noqa: F811
    db = async_db_session
    org_id, cookie = lumenvale
    r = await _build(client, db, org_id, cookie)
    pk = r["packages"]
    reclaim = next(p for p in pk if p["playbook"] == "licence-reclaim")
    reclaim["confirmed"] = {"low": 3, "high": 4}
    storage = next(p for p in pk if p["playbook"] == "data-storage")
    storage["included"] = False
    extra = rm.new_package("Client training", len(pk))
    out = await client.put(f"/orgs/{org_id}/roadmaps/{r['id']}", json={"packages": pk + [extra]}, cookies=cookie)
    assert out.status_code == 200, out.text
    d = out.json()
    got = next(p for p in d["packages"] if p["playbook"] == "licence-reclaim")
    assert got["basis"] == "confirmed" and got["confirmed"]["by"] == "admin@firm.test" and got["inferred"]["low"] == 2.0
    assert got["price"] == {"low": 3600, "high": 4800}  # demo rate card, 1,200 a day
    t = d["totals"]
    assert t["confirmed"] == {"low": 3.0, "high": 4.0, "packages": 1} and t["excluded"] == 1
    assert t["unestimated"] == ["Client training"]
    # A demo roadmap takes no per-engagement rates.
    assert (await client.put(f"/orgs/{org_id}/roadmaps/{r['id']}", json={"rate_overrides": {"Salesforce consultant": 1000}},
                             cookies=cookie)).status_code == 400
    # Overrides apply over the rate card (a client org's roadmap).
    card = {"roles": [{"role": "Salesforce consultant", "dayRate": 1200}]}
    over = rm.compute(d["packages"], {}, rm.demo_playbooks(), card, {"Salesforce consultant": 1000}, {})
    assert next(x for x in over["rates"] if x["role"] == "Salesforce consultant") == {
        "role": "Salesforce consultant", "dayRate": 1000.0, "override": True}
    assert next(p for p in over["packages"] if p["playbook"] == "licence-reclaim")["price"] == {"low": 3000, "high": 4000}
    bad = await client.put(f"/orgs/{org_id}/roadmaps/{r['id']}",
                           json={"packages": [{**pk[0], "confirmed": {"low": 5, "high": 2}}]}, cookies=cookie)
    assert bad.status_code == 400

    # The workbook keeps inferred and confirmed in separate columns and labels the demo.
    x = await client.get(f"/orgs/{org_id}/roadmaps/{r['id']}/export.xlsx", cookies=cookie)
    wb = load_workbook(io.BytesIO(x.content))
    assert wb.sheetnames == ["Summary", "Roadmap", "Findings", "Rate card", "Playbooks"]
    assert wb["Summary"]["A1"].value == "Demo data"
    head = [c.value for c in wb["Roadmap"][1]]
    row = next(r for r in wb["Roadmap"].iter_rows(min_row=2, values_only=True) if r[1] == "Dormant licence reclaim")
    assert row[head.index("Inferred low")] == 2.0 and row[head.index("Confirmed low")] == 3.0
    assert row[head.index("Effort used")] == "confirmed"
    assert wb["Findings"].max_row == 19


@pytest.mark.asyncio
async def test_reviews_and_rerun(client, async_db_session, lumenvale):  # noqa: F811
    db = async_db_session
    org_id, cookie = lumenvale
    await OrgAnalyzerService(db, org_id).run(actor_email="test")
    # The two seat findings have their own keys; one marked not useful starts outside the plan.
    from app.domain.models import OrgFinding
    from app.services.evidence import review_key
    seats = (await db.execute(select(OrgFinding).where(OrgFinding.organization_id == org_id,
                                                       OrgFinding.code == "LICENSE_SEATS_UNUSED"))).scalars().all()
    keys = {review_key("finding", f) for f in seats}
    assert len(keys) == 2
    platform = next(f for f in seats if "Platform" in f.title)
    db.add(ItemReview(organization_id=org_id, item_type="finding", item_key=review_key("finding", platform),
                      module="Health Report", group="license_waste", status="not_useful", reviewer="admin@firm.test",
                      title=platform.title))
    await db.commit()
    r = (await client.post(f"/orgs/{org_id}/roadmaps", json={"name": "Phase 1"}, cookies=cookie)).json()
    assert [f["title"] for f in r["notInRoadmap"]] == [platform.title] and r["notInRoadmap"][0]["review"] == "not_useful"
    # A rerun makes new finding rows; the roadmap still finds them by key.
    await OrgAnalyzerService(db, org_id).run(actor_email="test")
    again = (await client.get(f"/orgs/{org_id}/roadmaps/{r['id']}", cookies=cookie)).json()
    assert all(f["present"] for p in again["packages"] for f in p["findings"])
    assert again["totals"]["savingsCents"] == 18_318_000 - 150_000


@pytest.mark.asyncio
async def test_firm_data_and_roles(client, async_db_session):  # noqa: F811
    db = async_db_session
    org = Organization(name="Client Co", is_demo=False)
    db.add(org)
    await db.flush()
    admin = OrgUser(organization_id="system-org", email="boss@firm.test", role=OrgUserRole.ORG_ADMIN, password_hash="x",
                    is_active=True, is_email_verified=True)
    analyst = OrgUser(organization_id="system-org", email="analyst@firm.test", role=OrgUserRole.ANALYST, password_hash="x",
                      is_active=True, is_email_verified=True)
    db.add_all([admin, analyst])
    await db.flush()
    db.add(OrgAccessGrant(org_user_id=analyst.id, organization_id=org.id))
    await db.commit()
    ac, an = _cookie(admin), _cookie(analyst)
    # No Health Report yet.
    assert (await client.post(f"/orgs/{org.id}/roadmaps", json={}, cookies=an)).status_code == 409
    # A firm starts with nothing: no rates are ever invented.
    assert (await client.get("/firm/rate-card", cookies=an)).json()["roles"] == []
    assert (await client.get("/firm/playbooks", cookies=an)).json()["playbooks"] == []
    assert (await client.put("/firm/rate-card", json={"roles": [{"role": "Consultant", "dayRate": 900}]},
                             cookies=an)).status_code == 403
    assert (await client.put("/firm/rate-card", json={"roles": [{"role": "Consultant", "dayRate": -1}]},
                             cookies=ac)).status_code == 400
    card = (await client.put("/firm/rate-card", json={"currency": "gbp", "roles": [{"role": "Salesforce consultant",
                                                                                   "dayRate": 900}]}, cookies=ac)).json()
    assert card["currency"] == "GBP" and card["updatedBy"] == "boss@firm.test"
    ex = (await client.post("/firm/playbooks/examples", cookies=ac)).json()
    assert len(ex["playbooks"]) == 5 and all(p["example"] for p in ex["playbooks"])
    assert (await client.post("/firm/playbooks/examples", cookies=ac)).status_code == 409
    pbs = ex["playbooks"]
    pbs[0] = {**pbs[0], "baseDays": [1, 1.5], "example": False}
    saved = (await client.put("/firm/playbooks", json={"playbooks": pbs}, cookies=ac)).json()
    assert saved["version"] == 2 and saved["playbooks"][0]["version"] == "2026.10.2"
    bad = [{**pbs[0], "roles": {"Salesforce consultant": 0.5}}]
    assert (await client.put("/firm/playbooks", json={"playbooks": bad}, cookies=ac)).status_code == 400
    rows = (await db.execute(select(FirmSetting))).scalars().all()
    assert {(r.firm_id, r.key) for r in rows} == {("system-org", "rate_card"), ("system-org", "playbooks")}


@pytest.mark.asyncio
async def test_notes_masked_and_viewer_reads(client, async_db_session, lumenvale):  # noqa: F811
    db = async_db_session
    org_id, cookie = lumenvale
    r = await _build(client, db, org_id, cookie)
    from app.domain.models import UserSnapshot
    person = (await db.execute(select(UserSnapshot).where(UserSnapshot.organization_id == org_id).limit(1))).scalar_one()
    name = person.name
    pk = r["packages"]
    pk[0]["notes"] = f"{name} owns the approvals"
    assert (await client.put(f"/orgs/{org_id}/roadmaps/{r['id']}", json={"packages": pk}, cookies=cookie)).status_code == 200
    viewer = OrgUser(organization_id="system-org", email="viewer@firm.test", role=OrgUserRole.VIEWER, password_hash="x",
                     is_active=True, is_email_verified=True)
    db.add(viewer)
    await db.flush()
    db.add(OrgAccessGrant(org_user_id=viewer.id, organization_id=org_id))
    await db.commit()
    vc = _cookie(viewer)
    assert (await client.get(f"/orgs/{org_id}/roadmaps/{r['id']}", cookies=vc)).status_code == 200
    assert (await client.put(f"/orgs/{org_id}/roadmaps/{r['id']}", json={"name": "x"}, cookies=vc)).status_code == 403
    org = await db.get(Organization, org_id)
    privacy_mode.set_privacy(org, privacy_mode.MASKED, False, "test")
    await enforce_privacy_mode(db, org)
    await db.commit()
    got = (await client.get(f"/orgs/{org_id}/roadmaps/{r['id']}", cookies=cookie)).json()
    notes = next(p["notes"] for p in got["packages"] if p["notes"])
    assert name not in notes and "User " in notes
    assert (await client.delete(f"/orgs/{org_id}/roadmaps/{r['id']}", cookies=cookie)).status_code == 200


def test_phases_break_cycles_and_skip_excluded():
    a, b, c = (rm.new_package(t, i) for i, t in enumerate("abc"))
    a["after"], b["after"], c["after"] = [b["id"]], [a["id"]], [b["id"]]
    ph = rm.phases([a, b, c])
    assert ph[a["id"]] == 2 and ph[b["id"]] == 1 and ph[c["id"]] == 2
    b["included"] = False
    assert rm.phases([a, b, c])[c["id"]] == 1


@pytest.mark.asyncio
async def test_demo_previews_firm_rates_on_screen_only(client, async_db_session, lumenvale):  # noqa: F811
    db = async_db_session
    org_id, cookie = lumenvale
    r = await _build(client, db, org_id, cookie)
    assert r["rateSource"] == "demo" and r["firmRatesAvailable"] is False
    db.add(FirmSetting(firm_id="system-org", key="rate_card", value={"currency": "USD", "roles": [
        {"role": "Salesforce consultant", "dayRate": 950}]}))
    await db.commit()
    pv = (await client.get(f"/orgs/{org_id}/roadmaps/{r['id']}?rates=firm", cookies=cookie)).json()
    assert pv["rateSource"] == "firm-preview" and "on this screen only" in pv["labels"]["demo"]
    reclaim = next(p for p in pv["packages"] if p["playbook"] == "licence-reclaim")
    assert reclaim["blendedRate"] == 950 and reclaim["price"] == {"low": 1900, "high": 4750}
    assert next(p for p in pv["packages"] if p["playbook"] == "apex-coverage")["missingRates"]  # the firm has no developer rate
    plain = (await client.get(f"/orgs/{org_id}/roadmaps/{r['id']}", cookies=cookie)).json()
    assert plain["rateSource"] == "demo" and plain["firmRatesAvailable"] is True
    # Exports ignore the preview: the demo rate card only.
    x = await client.get(f"/orgs/{org_id}/roadmaps/{r['id']}/export.xlsx?rates=firm", cookies=cookie)
    rates = [row for row in load_workbook(io.BytesIO(x.content))["Rate card"].iter_rows(min_row=2, values_only=True)]
    assert all(src == "Demo rate card" for _, _, src in rates) and 950 not in [rate for _, rate, _ in rates]
