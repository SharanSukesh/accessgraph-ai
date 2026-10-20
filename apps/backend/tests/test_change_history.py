"""Change history (item 14, parity #25): Change-Risk Radar runs on the demo
org from the fixture's Setup Audit Trail, and the workbook's Change History
sheet keeps every pulled event once, across pulls."""
import io

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.domain.models import ChangeAuditEvent
from app.services.change_risk_radar import ChangeRiskRadarService
from tests.test_demo_orgs import client  # noqa: F401
from tests.test_demo_path import lumenvale  # noqa: F401
from tests.test_documentation_export import _rows


@pytest.mark.asyncio
async def test_change_risk_runs_on_the_demo_org(async_db_session, lumenvale):  # noqa: F811
    org_id, _ = lumenvale
    run = await ChangeRiskRadarService(async_db_session, org_id).run(actor_email="test")
    assert run.events_ingested == 60 and run.unique_actors == 8
    events = (await async_db_session.execute(select(ChangeAuditEvent).where(ChangeAuditEvent.run_id == run.id))).scalars().all()
    app = next(e for e in events if "DataBridge Sync" in e.display)
    assert app.section == "Connected Apps" and app.blast_tier == "critical"  # 85, plus 10 for an install
    delegated = next(e for e in events if e.section == "Sharing Defaults")
    assert delegated.delegate_user == "ayla.zielinski@lumenvale.example.com"
    # A shorter window returns only the recent part of the trail.
    week = await ChangeRiskRadarService(async_db_session, org_id, since_days=7).run(actor_email="test")
    assert 0 < week.events_ingested < 60


@pytest.mark.asyncio
async def test_change_history_sheet_keeps_each_event_once(client, async_db_session, lumenvale):  # noqa: F811
    org_id, cookie = lumenvale
    wb = load_workbook(io.BytesIO((await client.get(f"/orgs/{org_id}/documentation.xlsx", cookies=cookie)).content))
    assert _rows(wb["Change History"])[1][0]["When (UTC)"] == "Run Change-Risk Radar to pull the Setup Audit Trail."

    await ChangeRiskRadarService(async_db_session, org_id).run(actor_email="test")
    await ChangeRiskRadarService(async_db_session, org_id).run(actor_email="test")  # the same 60 events again
    wb = load_workbook(io.BytesIO((await client.get(f"/orgs/{org_id}/documentation.xlsx", cookies=cookie)).content))
    head, rows = _rows(wb["Change History"])
    assert len(rows) == 60 and rows == sorted(rows, key=lambda r: r["When (UTC)"], reverse=True)
    rule = next(r for r in rows if r["Change"] == "Created validation rule Old_Score_Range on Account")
    assert rule["Section"] == "Validation Rules" and rule["Changed by"] == "Idris Castellan"
    about = {row[0]: row[1] for row in wb["About"].iter_rows(values_only=True)}
    assert about["Change history"].startswith("60 Setup Audit Trail events from ")
    assert "from 2 Change-Risk Radar pulls" in about["Change history"] and "at least 180 days" in about["Change history"]
