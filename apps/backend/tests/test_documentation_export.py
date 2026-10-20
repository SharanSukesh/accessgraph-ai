"""Roadmap item 11, part a: the org documentation workbook."""
import io

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.domain.models import UserSnapshot
from app.services import privacy_mode
from app.services.effective_access import EffectiveAccessService
from app.services.report.documentation import NOT_AT_METADATA_ONLY, SHEETS, build_documentation_workbook
from tests.test_demo_orgs import client  # noqa: F401
from tests.test_demo_path import lumenvale  # noqa: F401
from tests.test_demo_privacy import ALIAS, _sync_demo


def _rows(ws):
    head = [c.value for c in ws[1]]
    return head, [dict(zip(head, r)) for r in ws.iter_rows(min_row=2, values_only=True)]


@pytest.mark.asyncio
async def test_org_wide_access_matches_the_per_user_views(async_db_session, lumenvale):  # noqa: F811
    db = async_db_session
    org_id, _ = lumenvale
    svc = EffectiveAccessService(db)
    every = await svc.org_access(org_id)
    users = (await db.execute(select(UserSnapshot.salesforce_id).where(
        UserSnapshot.organization_id == org_id, UserSnapshot.is_active.is_(True)))).scalars().all()
    assert len(every) == len(users)
    for uid in users:
        objs = {o["object"]: o["access"] for o in (await svc.get_user_object_access(org_id, uid))["objects"]}
        assert {k: v for k, v in every[uid]["objects"].items() if any(v.values())} == objs
        fields = {f"{f['objectName']}.{f['fieldName']}": f["access"]
                  for f in (await svc.get_user_field_access(org_id, uid))["fields"]}
        assert {k: v for k, v in every[uid]["fields"].items() if any(v.values())} == fields


@pytest.mark.asyncio
async def test_documentation_workbook_on_lumenvale(client, lumenvale):  # noqa: F811
    org_id, cookie = lumenvale
    r = await client.get(f"/orgs/{org_id}/documentation.xlsx", cookies=cookie)
    assert r.status_code == 200 and 'filename="demo-org-documentation-' in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == list(SHEETS)
    about = {row[0]: row[1] for row in wb["About"].iter_rows(values_only=True)}
    assert about["Demo data"] and about["Organization"] == "Lumenvale Systems (Demo)" and about["Privacy level"] == "Full"

    head, objects = _rows(wb["Objects"])
    by = {o["Object"]: o for o in objects}
    assert len(objects) == 29 and by["Invoice__c"]["Fields"] == 11
    assert by["Payment__c"]["Master-detail parent"] == "Invoice__c"
    assert "System Administrator (Profile)" in by["Account"]["Modify All granted by"]
    assert by["Account"]["Users with Read"] == 376

    _, fields = _rows(wb["Fields"])
    f = {(x["Object"], x["Field"]): x for x in fields}
    assert len(fields) == 262  # with item 13's Account.Old_Score__c and Fax_Ext__c
    assert f[("Employee__c", "National_ID__c")]["Encryption"] == "Shield"
    assert f[("Compensation__c", "Salary_Band__c")]["Encryption"] == "Shield"
    assert f[("Treasury_Account__c", "Account_Number__c")]["Type"] == "Text (Encrypted)(32)"
    assert (f[("Compensation__c", "Bonus__c")]["Encryption"], f[("Compensation__c", "Bonus__c")]["Required"]) == ("Not known", "Not known")
    assert f[("Invoice__c", "Days_Overdue__c")]["Formula"].startswith("IF(ISBLANK(")
    # Required fields follow the object; FLS fields count their own grants.
    assert f[("Account", "Name")]["Access follows"] == "Object access"
    assert f[("Account", "Name")]["Users with Read"] == 376
    assert f[("Employee__c", "National_ID__c")]["Access follows"] == "Field-level security"

    head, matrix = _rows(wb["Object Access"])
    assert head[1] == "Finance Analyst (Profile)" and any(h.endswith("(Group)") for h in head)
    account = next(m for m in matrix if m["Object"] == "Account")
    assert account["System Administrator (Profile)"] == "R C E D VA MA"
    _, unused = _rows(wb["Unused Layouts"])
    assert [(u["Object"], u["Layout"]) for u in unused] == [("Account", "Account - Legacy Layout"),
                                                             ("Case", "Case - Old Support Layout")]
    _, flows = _rows(wb["Flows"])
    discount = next(x for x in flows if x["API name"] == "Discount_Approval_Request")
    assert discount["Explicit access granted by"] == "Discount Approvals (in Sales Manager Bundle)"
    assert discount["Note"] == "Only through the explicit access listed."
    _, people = _rows(wb["Profile Users"])
    assert len(people) == 400 and all("@" not in (p["User"] or "") for p in people)


@pytest.mark.asyncio
async def test_masked_users_stay_masked(async_db_session):
    org = await _sync_demo(async_db_session, privacy_mode.MASKED)
    wb = load_workbook(io.BytesIO(await build_documentation_workbook(async_db_session, org)))
    _, people = _rows(wb["Profile Users"])
    _, assigned = _rows(wb["Assignments"])
    assert people and all(ALIAS.match(p["User"]) for p in people)
    assert assigned and all(ALIAS.match(a["User"]) for a in assigned)
    text = repr([[c for c in r] for ws in wb.worksheets for r in ws.iter_rows(values_only=True)])
    assert "Mina Tanaka" not in text and "lumenvale.example.com" not in text


@pytest.mark.asyncio
async def test_metadata_only_leaves_out_people(async_db_session):
    org = await _sync_demo(async_db_session, privacy_mode.METADATA_ONLY)
    wb = load_workbook(io.BytesIO(await build_documentation_workbook(async_db_session, org)))
    assert wb["Assignments"]["A2"].value == NOT_AT_METADATA_ONLY
    assert wb["Profile Users"]["A2"].value == NOT_AT_METADATA_ONLY
    head, objects = _rows(wb["Objects"])
    assert "Users with Read" not in head and len(objects) == 29
    _, fields = _rows(wb["Fields"])
    assert len(fields) == 262  # configuration is the same at every level
    head, rules = _rows(wb["Validation Rules"])
    assert len(rules) == 26 and next(r for r in rules if r["Rule"] == "Industry_Required")[
        "Custom permissions it checks"] == "Bypass_Validation"
    head, perms = _rows(wb["Custom Permissions"])
    assert "Active users granted it" not in head and len(perms) == 3


# --------------------------------------------------------------------------- part b

def test_sharing_keyword_is_all_that_is_kept():
    from app.services.org_docs import sharing_of
    assert sharing_of("public with sharing class A {}") == "with sharing"
    assert sharing_of("/* without sharing */ public inherited sharing class B {}") == "inherited sharing"
    assert sharing_of("// with sharing\npublic class C { class D without sharing {} }") == "not declared"
    assert sharing_of("@isTest private without sharing class E {}") == "without sharing"
    assert sharing_of(None) is None


def test_login_restrictions_from_profile_metadata():
    from app.services.ui_access import login_restrictions
    md = {"loginIpRanges": [{"startAddress": "192.0.2.0", "endAddress": "192.0.2.255", "description": "VPN"}],
          "loginHours": {"mondayStart": 360, "mondayEnd": 1260}}
    assert login_restrictions(md) == {"ip_ranges": [{"start": "192.0.2.0", "end": "192.0.2.255", "description": "VPN"}],
                                      "hours": {"monday": {"start": 360, "end": 1260}}}
    assert login_restrictions({"loginHours": [{"fridayStart": 0, "fridayEnd": 600}]}) == {"hours": {"friday": {"start": 0, "end": 600}}}
    assert login_restrictions({}) == {}


@pytest.mark.asyncio
async def test_live_documentation_reads_degrade_part_by_part():
    from app.services import org_docs

    queries = []

    class Client:
        async def query_tooling(self, q):
            queries.append(q)
            if "FROM CustomObject" in q:
                return [{"Id": "01I000000000001AAA", "DeveloperName": "Invoice", "NamespacePrefix": None}]
            if q.startswith("SELECT Metadata FROM ValidationRule"):
                return [{"Metadata": {"errorConditionFormula": "Amount__c <= 0"}}]
            if "FROM ValidationRule" in q:
                return [{"Id": "03d1", "ValidationName": "Amount_Positive", "Active": True, "ErrorMessage": "Positive.",
                         "EntityDefinition": {"QualifiedApiName": "Invoice__c"}}]
            if "FROM ApexClass" in q:
                return [{"Id": "01p1", "Name": "A", "ApiVersion": 58.0, "Status": "Active", "IsValid": True,
                         "LengthWithoutComments": 10, "Body": "public without sharing class A {}"}]
            if "FROM ApexTrigger" in q:
                return [{"Id": "01q1", "Name": "T", "TableEnumOrId": "01I000000000001AAA", "UsageBeforeInsert": True}]
            raise RuntimeError("not allowed")

        async def query_all(self, q):
            queries.append(q)
            if "FROM PermissionSetTabSetting" in q:
                return [{"ParentId": "0PS1", "Name": "standard-Account", "Visibility": "DefaultOn"}]
            raise RuntimeError("not allowed")

        async def describe_layouts(self, name):
            return {"recordTypeMappings": [{"name": "Renewal", "picklistsForRecordType": [
                {"picklistName": "StageName", "picklistValues": [{"value": "Closed Won"}]}]}]}

        async def get_code_coverage(self):
            return {"items": [{"name": "A", "type": "class", "covered": 3, "uncovered": 1, "is_test": False}]}

    out = await org_docs.collect_live(Client(), ["Opportunity"])
    assert out["validation_rules"][0]["formula"] == "Amount__c <= 0"
    assert out["apex_classes"][0]["sharing"] == "without sharing" and "Body" not in repr(out)
    assert out["apex_triggers"][0]["object"] == "Invoice__c" and out["apex_triggers"][0]["events"] == ["before insert"]
    assert out["record_type_picklists"] == {"Opportunity": {"Renewal": {"StageName": ["Closed Won"]}}}
    assert out["tab_settings"] == [{"parent_id": "0PS1", "tab": "standard-Account", "visibility": "DefaultOn"}]
    assert "Visualforce pages couldn't be read." in out["notes"] and "static_resources" not in out
    assert all(q.lstrip().upper().startswith("SELECT") for q in queries)  # reads only


@pytest.mark.asyncio
async def test_part_b_sheets_on_lumenvale(client, lumenvale):  # noqa: F811
    org_id, cookie = lumenvale
    wb = load_workbook(io.BytesIO((await client.get(f"/orgs/{org_id}/documentation.xlsx", cookies=cookie)).content))
    _, rules = _rows(wb["Validation Rules"])
    assert len(rules) == 26 and sum(r["Active"] == "No" for r in rules) == 3  # 26th: item 13's Old_Score_Range
    industry = next(r for r in rules if r["Rule"] == "Industry_Required")
    assert industry["Custom permissions it checks"] == "Bypass_Validation (9 active users hold it)"
    _, rt = _rows(wb["Picklists by Record Type"])
    billing = next(r for r in rt if r["Record type"] == "Billing Inquiry" and r["Field"] == "Status")
    assert (billing["Values available"], billing["Values left out"]) == ("New, Working, Closed", "Escalated")
    head, tabs = _rows(wb["Tab Visibility"])
    comp = next(t for t in tabs if t["Tab"] == "Compensation__c")
    assert comp["Compensation Admin"] == "Default On" and comp["System Administrator (Profile)"] == "Default Off"
    _, logins = _rows(wb["Login Restrictions"])
    fd = next(r for r in logins if r["Profile"] == "Finance Director")
    assert fd["Login IP ranges"].startswith("198.51.100.0 to 198.51.100.63 (Finance floor)")
    assert next(r for r in logins if r["Profile"] == "Support Agent")["Login hours"].startswith("Mon 06:00 to 21:00")
    _, apex = _rows(wb["Apex Classes"])
    acct = next(a for a in apex if a["Class"] == "AccountHandler")
    assert (acct["Sharing"], acct["Coverage %"]) == ("without sharing", 75.6)
    assert len(apex) == 81 and len(_rows(wb["Apex Triggers"])[1]) == 12
    _, vf = _rows(wb["Visualforce"])
    legacy = next(v for v in vf if v["Name"] == "LegacyTreasuryDashboard")
    assert (legacy["API version"], legacy["Access granted by"], legacy["Active users granted it"]) == (29, "Treasury Data Steward", 1)
    _, apps = _rows(wb["Apps"])
    assert next(a for a in apps if a["App"] == "Lumenvale Sales")["Active users granted it"] == 161
    for sheet, n in (("Approval Processes", 3), ("Buttons and Links", 5), ("List Views", 13), ("Field Sets", 2),
                     ("Lightning Components", 7), ("Static Resources", 4), ("Custom Permissions", 3)):
        assert len(_rows(wb[sheet])[1]) == n, sheet
