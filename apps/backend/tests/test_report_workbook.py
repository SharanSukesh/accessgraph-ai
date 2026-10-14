"""The Health Report Excel export, built from the synthetic demo org."""
import io

import pytest
from openpyxl import load_workbook

from app.domain.models import Organization, SyncJob, SyncStatus
from app.ingestion.orchestrator import SyncOrchestrator
from app.services.license_fit import LicenseFitService
from app.services.org_analyzer import OrgAnalyzerService
from tests.test_demo_orgs import _cookie, _users, client  # noqa: F401  (fixture)

SHEETS = ["Summary", "Findings", "License Fit", "Anomalies", "User Risk", "Sprawl Summary", "Data Quality"]


@pytest.mark.asyncio
async def test_report_workbook_has_every_sheet_with_rows(client, async_db_session):  # noqa: F811
    admin, _ = await _users(async_db_session)
    org = Organization(name="Lumenvale Systems (Demo)", is_demo=True,
                       settings={"demo_fixture": "synthetic_mid_market"})
    async_db_session.add(org)
    await async_db_session.flush()
    job = SyncJob(organization_id=org.id, status=SyncStatus.PENDING)
    async_db_session.add(job)
    await async_db_session.commit()
    await SyncOrchestrator(async_db_session).run_sync(org.id, job.id)
    await OrgAnalyzerService(async_db_session, org.id).run(actor_email="test")
    await LicenseFitService(async_db_session, org.id).run(actor_email="test")

    r = await client.get(f"/orgs/{org.id}/org-analyzer/report.xlsx", cookies=_cookie(admin))

    assert r.status_code == 200
    assert 'filename="demo-org-assessment-' in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == SHEETS
    for name in SHEETS:
        assert wb[name].max_row >= 2, name

    summary = {row[0].value: row[1].value for row in wb["Summary"].iter_rows(max_col=2)}
    assert summary["Organization"] == "Lumenvale Systems (Demo)"
    assert "Demo data" in summary
    assert wb["Findings"].freeze_panes == "A2" and wb["Findings"].auto_filter.ref
    assert wb["Anomalies"].max_row - 1 == 8  # top 2% of 400 users
    assert wb["Anomalies"]["A2"].value == 1
    assert wb["License Fit"].max_row - 1 == 400
    savings_col = next(c.column for c in wb["Findings"][1] if c.value == "Est. annual savings")
    assert wb["Findings"].cell(row=2, column=savings_col).number_format.startswith('"$"')
    # Live-only modules never ran on a demo org: the sheet says so.
    assert "not been run" in wb["Data Quality"]["A2"].value
